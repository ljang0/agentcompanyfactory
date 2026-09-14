import io
import re
import socket
import subprocess
import tarfile
from types import SimpleNamespace
from typing import ClassVar

import pytest

from company_envs.storage import read, write
from company_envs.world import hub_vm


@pytest.fixture
def company(tmp_path):
    folder = tmp_path / "company"
    write(
        folder / "tasks/task/workflow.json",
        {"worker_ids": ["boss", "analyst"], "manager_id": "boss", "private": "SECRET-DESIGN"},
    )
    write(
        folder / "tasks/task/assignment.json",
        {
            "title": "Reconcile deliveries",
            "brief": "Resolve the delayed delivery.",
            "private_extra": "DO-NOT-COPY",
        },
    )
    write(
        folder / "company.json",
        {"name": "Harbor Logistics", "workers": [{"id": "analyst", "title": "Delivery Analyst"}]},
    )
    write(
        folder / "world/identities.json",
        {"analyst": {"chat": {"userId": "analyst", "fullName": "Dana Reyes"}}},
    )
    write(folder / "world/worker_apps.json", {"boss": ["chat", "sheets"], "analyst": ["chat"]})
    write(folder / "runtime/sessions.json", {"sid": "company-ep1", "workers": ["boss", "analyst"]})
    write(
        folder / "runtime/endpoints.json",
        {
            "apps": {
                "chat": {
                    "harness_url": "http://127.0.0.1:8000",
                    "workers": {
                        "boss": "http://0.0.0.0:8100/?sid=company-ep1",
                        "analyst": "http://0.0.0.0:8101/?sid=company-ep1",
                    },
                },
                "sheets": {
                    "harness_url": "http://127.0.0.1:8001",
                    "workers": {
                        "boss": "http://0.0.0.0:8200/?sid=company-ep1",
                        "analyst": "http://0.0.0.0:8201/?sid=company-ep1",
                    },
                },
            }
        },
    )
    for worker in ("boss", "analyst"):
        root = folder / "world/materials" / worker
        (root / "notes").mkdir(parents=True)
        (root / "notes/onboarding.md").write_text(f"Welcome {worker}")
        (root / "exports").mkdir()
        (root / "exports/open orders.csv").write_bytes(b"id,status\n1,open\n")
        (root / "Documents/Policies").mkdir(parents=True)
        (root / "Documents/Policies/Rate card.md").write_text("# Rates\n\nStandard tariff.")
    return folder


def launch(folder, **kwargs):
    return hub_vm.launch_company(
        folder,
        "task",
        base_image=folder / "base.qcow2",
        browser_dir=folder / "browser",
        host_ip="192.168.1.20",
        dry_run=True,
        **kwargs,
    )


@pytest.fixture(autouse=True)
def _snapshots_on(monkeypatch):
    # Snapshot saving is opt-in in production (it writes guest RAM into the overlay); these
    # tests exercise the snapshot path, so they turn it on.
    monkeypatch.setattr(hub_vm, "SAVE_SNAPSHOT", True)


def test_dry_run_delivery_and_private_boundaries(company, monkeypatch):
    def no_vm(*args, **kwargs):
        pytest.fail("Dry run must never instantiate QEMU")

    monkeypatch.setattr(hub_vm, "HubWorkerVM", no_vm)
    monkeypatch.setattr(hub_vm.subprocess, "run", no_vm)
    report = launch(company)
    root = company / "runtime/vms"
    assert report == read(root / "LAUNCH.json")
    assert report["workers"] == ["boss", "analyst"]
    assert report["boss"] == "boss"
    assert report["started_at"]
    assert report["dry_run"] is True
    assert report["boss_assumption"] is None
    for worker in report["workers"]:
        target = root / worker
        vm = read(target / "vm.json")
        assert vm["worker"] == worker
        assert vm["boss"] == (worker == "boss")
        assert vm["dry_run"] is True
        assert vm["vm"]["process"] is None
        desktop = target / "guest/Desktop"
        assert (desktop / "notes/onboarding.md").read_text() == f"Welcome {worker}"
        assert (desktop / "exports/open orders.csv").read_bytes() == b"id,status\n1,open\n"
        # Home folders named by a material land as written; everything else lands on the Desktop.
        assert (target / "guest/Documents/Policies/Rate card.md").read_text().startswith("# Rates")
        assert not (desktop / "Documents").exists()
        assert vm["materials_delivered"] == [
            "/home/ga/Desktop/exports/open orders.csv",
            "/home/ga/Desktop/notes/onboarding.md",
            "/home/ga/Documents/Policies/Rate card.md",
        ]
        assert (desktop / "ASSIGNMENT.md").exists() == (worker == "boss")
        assert (desktop / "ROLE.md").exists() == (worker != "boss")
        bookmarks = read(target / "guest/worker-browser/Default/Bookmarks")
        assert {b["url"] for b in bookmarks["roots"]["bookmark_bar"]["children"]} == set(vm["urls"].values())
        assert {b["name"] for b in bookmarks["roots"]["bookmark_bar"]["children"]} <= {"Chat", "Sheets"}
        preferences = read(target / "guest/worker-browser/Default/Preferences")
        assert preferences["session"] == {
            "restore_on_startup": 4,
            "startup_urls": [hub_vm.START_PAGE, *vm["urls"].values()],
        }
        assert preferences["bookmark_bar"] == {"show_on_all_tabs": True}
        page = (desktop / "APPS.html").read_text()
        assert "Harbor Logistics" in page and "<b>Chat</b>" in page
        assert "http" not in re.sub(r'href="[^"]*"', "", page)  # addresses only behind the tiles
        for word in ("mock", "proxy", "sid", "worker", "harness"):
            assert word not in re.sub(r'href="[^"]*"', "", page).lower()
        note = (desktop / ("ASSIGNMENT.md" if worker == "boss" else "ROLE.md")).read_text()
        assert "## Your apps" in note and "- Chat \u2014 Company application" in note
        assert "open as browser tabs" in note and "http" not in note
        for path in target.rglob("*"):
            if path.is_file():
                content = path.read_text()
                for forbidden in (
                    "harness_url",
                    "127.0.0.1:800",
                    "192.168.1.20:800",
                    "SECRET-DESIGN",
                    "DO-NOT-COPY",
                    "0.0.0.0",
                    "10.0.2.100",
                ):
                    assert forbidden not in content
                if worker == "analyst":
                    for forbidden in (":8100", ":8200", ":8201", "Resolve the delayed delivery"):
                        assert forbidden not in content
                    assert "sheets.harbor-logistics.internal" not in content
                    assert "10.0.2.102" not in content
    assert "Resolve the delayed delivery." in (root / "boss/guest/Desktop/ASSIGNMENT.md").read_text()
    assert "- Sheets \u2014 Company application" in (root / "boss/guest/Desktop/ASSIGNMENT.md").read_text()
    assert "Delivery Analyst" in (root / "analyst/guest/Desktop/ROLE.md").read_text()
    assert "Await delegation" in (root / "analyst/guest/Desktop/ROLE.md").read_text()
    assert "Dana Reyes \u00b7 Delivery Analyst" in (root / "analyst/guest/Desktop/APPS.html").read_text()
    # No identity and no company title: the onboarding heading stands in for both.
    assert "Signed in as Welcome boss</p>" in (root / "boss/guest/Desktop/APPS.html").read_text()
    assert read(root / "analyst/vm.json")["urls"] == {
        "chat": "http://chat.harbor-logistics.internal/?sid=company-ep1"
    }
    assert read(root / "analyst/vm.json")["addresses"] == {
        "chat": {"host": "chat.harbor-logistics.internal", "ip": "10.0.2.101", "proxy_port": 8101}
    }
    assert read(root / "boss/vm.json")["addresses"] == {
        "chat": {"host": "chat.harbor-logistics.internal", "ip": "10.0.2.101", "proxy_port": 8100},
        "sheets": {"host": "sheets.harbor-logistics.internal", "ip": "10.0.2.102", "proxy_port": 8200},
    }


def test_stop_and_reset_restore_files_and_discard_overlay(company):
    first = launch(company)
    root = company / "runtime/vms"
    (root / "analyst/worker.qcow2").write_text("old overlay")
    (root / "boss/guest/Desktop/ASSIGNMENT.md").unlink()
    (root / "analyst/guest/Desktop/stale.txt").write_text("old guest work")
    assert hub_vm.stop_company(company)["status"] == "stopped"
    assert (root / "analyst/worker.qcow2").exists()
    second = hub_vm.reset_company(company)
    assert second["inputs"] == first["inputs"]
    assert second["started_at"] != first["started_at"]
    assert (root / "boss/guest/Desktop/ASSIGNMENT.md").is_file()
    assert not (root / "analyst/guest/Desktop/ASSIGNMENT.md").exists()
    assert not (root / "analyst/guest/Desktop/stale.txt").exists()
    assert not (root / "analyst/worker.qcow2").exists()


def test_missing_endpoints(company):
    (company / "runtime/endpoints.json").unlink()
    with pytest.raises(FileNotFoundError, match="hub-serve"):
        launch(company)
    assert not (company / "runtime/vms").exists()


def test_default_roster_includes_worker_apps_and_boss_fallback(company):
    write(company / "tasks/task/workflow.json", {"worker_ids": ["boss"], "manager_id": None})
    report = launch(company)
    assert report["workers"] == ["boss", "analyst"]
    assert report["boss"] == "boss"
    assert "first task worker" in report["boss_assumption"]


def test_explicit_subset_and_manager_not_first(company):
    write(company / "tasks/task/workflow.json", {"worker_ids": ["analyst", "boss"], "manager_id": "boss"})
    assert launch(company, workers=["boss"])["workers"] == ["boss"]
    assert not (company / "runtime/vms/analyst").exists()


@pytest.mark.parametrize("workers", [["analyst"], ["boss", "unknown"], [], "boss"])
def test_invalid_subset(company, workers):
    with pytest.raises((ValueError, TypeError), match="workers"):
        launch(company, workers=workers)


@pytest.mark.parametrize(
    "bad",
    [
        "http://127.0.0.1:8000",
        "http://192.168.1.20:8000",
        "harness_url",
        "http://0.0.0.0:8100/?sid=company-ep1",
        "http://192.168.1.20:8201/?sid=company-ep1",
        "sheets.harbor-logistics.internal",  # another worker's app by guest name
        "http://10.0.2.102/",  # another worker's app by guest address
        "10.0.2.100:8101",  # the retired shared guest address
    ],
)
def test_materials_cannot_leak_private_or_unassigned_urls(company, bad):
    (company / "world/materials/analyst/notes/onboarding.md").write_text(bad)
    with pytest.raises(ValueError, match="private/unassigned"):
        launch(company)
    assert not (company / "runtime/vms").exists()


def test_materials_may_name_the_workers_own_app_address(company):
    (company / "world/materials/analyst/notes/onboarding.md").write_text(
        "Chat: http://chat.harbor-logistics.internal/ (10.0.2.101); the gateway is 10.0.2.2."
    )
    launch(company)


def test_identity_names_are_vm_bound_content(company):
    write(
        company / "world/identities.json",
        {"analyst": {"chat": {"fullName": "http://192.168.1.20:8201/?sid=company-ep1"}}},
    )
    with pytest.raises(ValueError, match="private/unassigned"):
        launch(company)


@pytest.mark.parametrize(
    "name, expected",
    [
        ("notes/onboarding.md", "Desktop/notes/onboarding.md"),
        ("Desktop/Closed orders.xlsx", "Desktop/Closed orders.xlsx"),
        ("Documents/Policies/Rates.pdf", "Documents/Policies/Rates.pdf"),
        ("Downloads/scan_0021.pdf", "Downloads/scan_0021.pdf"),
        ("Pictures/badge.png", "Pictures/badge.png"),
        ("Music/hold.mp3", "Music/hold.mp3"),
        ("documents/lowercase.md", "Desktop/documents/lowercase.md"),
    ],
)
def test_materials_land_in_the_home_folder_they_name(name, expected):
    from company_envs.world.documents import guest_path

    assert guest_path(name) == expected


def test_materials_that_collide_or_shadow_launcher_files_are_rejected(company):
    root = company / "world/materials/analyst"
    (root / "Desktop").mkdir()
    (root / "Desktop/notes").mkdir()
    (root / "Desktop/notes/onboarding.md").write_text("duplicate")
    with pytest.raises(ValueError, match="collide"):
        launch(company)
    (root / "Desktop/notes/onboarding.md").unlink()
    (root / "Desktop/ROLE.md").write_text("mine")
    with pytest.raises(ValueError, match="reserved launcher"):
        launch(company)


def test_app_labels_and_guest_addresses():
    assert hub_vm.app_label("gmail_mock") == ("Gmail", "Email")
    assert hub_vm.app_label("google_calendar_mock") == ("Calendar", "Meetings and schedules")
    assert hub_vm.app_label("Zendesk_mock") == ("Zendesk", "Company application")
    assert hub_vm.app_label("outlook_web_mock")[0] == "Outlook"
    assert hub_vm.host_label("google_sheets_mock") == "sheets"
    assert hub_vm.host_label("PACS-viewer_mock") == "pacsviewer"
    assert hub_vm.company_slug({"name": "Hearthway Furniture Rental"}) == "hearthway-furniture-rental"
    assert hub_vm.company_slug({"name": "  ", "id": "cort"}) == "cort"
    assert hub_vm.company_slug({}) == "company"
    plan = hub_vm.address_plan(["gmail_mock", "slack_mock", "tableau_mock"], "hearthway")
    assert plan == {
        "gmail_mock": {"host": "gmail.hearthway.internal", "ip": "10.0.2.101"},
        "slack_mock": {"host": "slack.hearthway.internal", "ip": "10.0.2.102"},
        "tableau_mock": {"host": "tableau.hearthway.internal", "ip": "10.0.2.103"},
    }
    assert hub_vm.hosts_lines(plan) == (
        "10.0.2.101 gmail.hearthway.internal\n10.0.2.102 slack.hearthway.internal\n"
        "10.0.2.103 tableau.hearthway.internal\n"
    )
    with pytest.raises(ValueError, match="share the guest hostname"):
        hub_vm.address_plan(["gmail_mock", "gmail"], "hearthway")
    with pytest.raises(ValueError, match="Too many apps"):
        hub_vm.address_plan([f"app{i}_mock" for i in range(160)], "hearthway")


def test_start_page_is_a_company_launcher_without_raw_addresses():
    urls = {
        "gmail_mock": "http://gmail.hw.internal/?sid=s&x=<1>",
        "slack_mock": "http://slack.hw.internal/?sid=s",
    }
    page = hub_vm.start_page("Hearthway & Co", "Della Cartwright", "Plant Manager", urls).decode()
    assert page.startswith("<!doctype html>") and "<style>" in page and "<script" not in page
    assert "Hearthway &amp; Co" in page and "Della Cartwright \u00b7 Plant Manager" in page
    assert (
        '<a class="tile" href="http://gmail.hw.internal/?sid=s&amp;x=&lt;1&gt;"><b>Gmail</b><span>Email</span></a>'
        in page
    )
    assert "<b>Slack</b><span>Team chat</span>" in page
    assert "http" not in re.sub(r'href="[^"]*"', "", page)
    assert len(page) < 4000
    assert hub_vm.apps_note(["gmail_mock", "slack_mock"]) == (
        "## Your apps\n\n- Gmail \u2014 Email\n- Slack \u2014 Team chat\n\n"
        "They are open as browser tabs and pinned in the Company apps bookmarks bar.\n"
    )


def test_symlink_materials_rejected(company):
    (company / "world/materials/analyst/private").symlink_to(company / "tasks/task/workflow.json")
    with pytest.raises(ValueError, match="regular files"):
        launch(company)


def test_missing_allowed_endpoint_does_not_fall_back_to_harness(company):
    endpoints = read(company / "runtime/endpoints.json")
    del endpoints["apps"]["chat"]["workers"]["analyst"]
    write(company / "runtime/endpoints.json", endpoints)
    with pytest.raises(ValueError, match="Missing proxy endpoint"):
        launch(company)


def test_existing_launch_requires_reset(company):
    launch(company)
    with pytest.raises(FileExistsError, match="reset_company"):
        launch(company)


class FakeVM:
    instances: ClassVar[list] = []
    fail = False

    def __init__(self, directory, worker, base_image, clock=None):
        self.directory = directory / worker
        self.directory.mkdir()
        (self.directory / "worker.qcow2").write_text("overlay")
        self.worker = worker
        self.command = ["qemu", "-netdev", "user,id=n0,hostfwd=tcp:127.0.0.1:12345-:22"]
        self.ssh_port, self.cdp_port = 12345, 12346
        self.process = self.tunnel = None
        self.uploads, self.commands = [], []
        self.closed = self.opened = False
        self.instances.append(self)

    def boot(self, *, snapshot=False):
        self.restored = snapshot
        self.process = SimpleNamespace(pid=123)
        if self.fail and self.worker == "analyst":
            raise RuntimeError("boot failed")

    def upload(self, source, destination):
        with tarfile.open(source) as archive:
            self.uploads.append((destination, archive.getnames()))

    def run(self, command, **kwargs):
        self.commands.append(command)
        return SimpleNamespace(returncode=0)

    def open_browser(self, urls=(), hosts=None):
        self.opened = True
        self.tabs, self.hosts = list(urls), hosts
        self.tunnel = SimpleNamespace(pid=456)

    def save_snapshot(self):
        assert self.opened
        self.saved = True

    @classmethod
    def from_snapshot(cls, directory, worker, record):
        vm = cls.__new__(cls)
        vm.directory, vm.worker = directory / worker, worker
        vm.command = record["command"]
        vm.ssh_port, vm.cdp_port = record["ssh_port"], record["cdp_port"]
        vm.process = vm.tunnel = None
        vm.uploads, vm.commands = [], []
        vm.closed = vm.opened = vm.saved = False
        cls.instances.append(vm)
        return vm

    def connect_browser(self):
        self.tunnel = SimpleNamespace(pid=789)

    def close(self):
        self.closed = True


@pytest.fixture
def fake_vm(monkeypatch, company):
    FakeVM.instances = []
    FakeVM.fail = False
    monkeypatch.setattr(hub_vm, "HubWorkerVM", FakeVM)
    monkeypatch.setattr(hub_vm, "_preflight", lambda *args: None)
    monkeypatch.setattr(hub_vm, "_process_identity", lambda pid: {"pid": pid, "start_ticks": "test"})
    (company / "browser").mkdir()
    (company / "browser/chrome").write_text("browser double")
    return FakeVM


def real_branch(company):
    return hub_vm.launch_company(
        company,
        "task",
        base_image=company / "base.qcow2",
        browser_dir=company / "browser",
        host_ip="192.168.1.20",
    )


def test_real_launch_orchestration_with_doubles(company, fake_vm, monkeypatch):
    assert real_branch(company)["status"] == "running"
    assert len(fake_vm.instances) == 2
    for vm in fake_vm.instances:
        assert vm.opened and vm.saved
        guest = dict(vm.uploads)["/home/ga/hub-guest.tar"]
        assert ("Desktop/ASSIGNMENT.md" in guest) == (vm.worker == "boss")
        assert "worker-browser/Default/Bookmarks" in guest
        assert "Desktop/exports/open orders.csv" in guest
        assert "Documents/Policies/Rate card.md" in guest and "Desktop/Documents" not in str(guest)
        assert not any("endpoints" in name or "workflow" in name for name in guest)
        receipt = read(vm.directory / "vm.json")
        assert receipt["vm"]["process"]["pid"] == 123
        assert receipt["vm"]["tunnel"]["pid"] == 456
        assert receipt["hosts_written"] is True
        hosts = next(command for command in vm.commands if "/etc/hosts" in command)
        assert "sudo -S" in hosts and "10.0.2.101 chat.harbor-logistics.internal" in hosts
        curls = [command for command in vm.commands if "curl --fail" in command]
        assert any("--resolve chat.harbor-logistics.internal:80:10.0.2.101 " in command for command in curls)
        assert curls and all("10.0.2.100" not in command for command in curls)
        assert vm.tabs == list(receipt["urls"].values())
        assert vm.hosts == {a["host"]: a["ip"] for a in receipt["addresses"].values()}
    stopped = []
    monkeypatch.setattr(hub_vm, "_stop_process", stopped.append)
    assert hub_vm.stop_company(company)["status"] == "stopped"
    assert [item["pid"] for item in stopped] == [456, 123, 456, 123]
    assert all((vm.directory / "worker.qcow2").exists() for vm in fake_vm.instances)


def test_boot_failure_closes_every_owned_vm(company, fake_vm):
    fake_vm.fail = True
    with pytest.raises(RuntimeError, match="boot failed"):
        real_branch(company)
    assert all(vm.closed for vm in fake_vm.instances)
    assert read(company / "runtime/vms/LAUNCH.json")["status"] == "failed"
    assert read(company / "runtime/vms/analyst/vm.json")["vm"]["process"]["pid"] == 123


def test_stop_process_checks_birth_identity():
    # A harmless local child exercises cross-invocation cleanup without a VM.
    process = subprocess.Popen(["sleep", "30"])
    try:
        identity = hub_vm._process_identity(process.pid)
        hub_vm._stop_process(identity | {"start_ticks": "wrong"})
        assert process.poll() is None
        hub_vm._stop_process(identity)
        assert process.poll() is not None
        hub_vm._stop_process(identity)  # Already stopped is idempotent.
    finally:
        if process.poll() is None:
            process.kill()
        process.wait()


def test_hub_network_uses_nat_without_native_guest_forward(monkeypatch, tmp_path):
    seen = {}

    def init(self, directory, worker, base_image, app_port, clock=None):
        self.ssh_port = 1234
        seen["clock"] = clock
        self.command = ["qemu", "-netdev", "old-native-forward", "-device", "virtio-net-pci,netdev=n0"]

    monkeypatch.setattr(hub_vm.WorkerVM, "__init__", init)
    vm = hub_vm.HubWorkerVM(tmp_path, "boss", tmp_path / "base", clock="2026-09-08T09:00:00")
    assert vm.command[2] == "user,id=n0,restrict=on,ipv6=off,hostfwd=tcp:127.0.0.1:1234-:22"
    assert seen["clock"] == "2026-09-08T09:00:00"


def test_reference_instant_comes_from_the_seed(tmp_path):
    (tmp_path / "world").mkdir()
    assert hub_vm.reference_instant(tmp_path) is None
    (tmp_path / "world" / "SEED.json").write_text('{"reference_date": "2026-09-08"}')
    assert hub_vm.reference_instant(tmp_path) == "2026-09-08T09:00:00"


def test_cleanup_failure_does_not_skip_other_vms(company, fake_vm, monkeypatch):
    fake_vm.fail = True

    def close(vm):
        vm.closed = True
        if vm.worker == "boss":
            raise OSError("cleanup failed")

    monkeypatch.setattr(fake_vm, "close", close)
    with pytest.raises(RuntimeError, match="boot failed"):
        real_branch(company)
    assert all(vm.closed for vm in fake_vm.instances)
    assert read(company / "runtime/vms/LAUNCH.json")["cleanup_errors"] == ["boss: cleanup failed"]


def test_role_uses_onboarding_when_company_title_is_unavailable(company):
    (company / "company.json").unlink()
    launch(company)
    assert "Welcome analyst" in (company / "runtime/vms/analyst/guest/Desktop/ROLE.md").read_text()


@pytest.mark.parametrize("display", ["X0", ""])
def test_browser_requires_desktop_and_waits_for_homepage(monkeypatch, display):
    vm = object.__new__(hub_vm.HubWorkerVM)
    vm.worker, vm.cdp_port = "boss", 12345
    vm.ssh_base, vm.env = ["ssh", "ga@127.0.0.1"], {}
    commands = []

    def run(command):
        commands.append(command)
        return SimpleNamespace(stdout=display if command == "ls /tmp/.X11-unix" else "")

    vm.run = run
    popen_calls = []

    def popen(command, **kwargs):
        popen_calls.append(command)
        return SimpleNamespace(poll=lambda: None)

    monkeypatch.setattr(hub_vm.subprocess, "Popen", popen)
    responses = iter([b'[{"url":"about:blank"}]', b'[{"url":"file:///home/ga/Desktop/APPS.html"}]'])
    tabs = ["http://chat.hl.internal/?sid=1", "http://sheets.hl.internal/?sid=1"]
    hosts = {"chat.hl.internal": "10.0.2.101", "sheets.hl.internal": "10.0.2.102"}
    requests = []

    def urlopen(url, **kwargs):
        requests.append(url)
        return io.BytesIO(next(responses))

    monkeypatch.setattr(hub_vm, "urlopen", urlopen)
    monkeypatch.setattr(hub_vm.time, "sleep", lambda _: None)
    if display:
        assert vm.open_browser(tabs, hosts) == 12345
        assert len(requests) == 2
        launch_line = commands[-1]
        assert "--user-data-dir=/home/ga/worker-browser" in launch_line
        assert "--no-first-run --no-default-browser-check --test-type" in launch_line
        assert (
            "'--host-resolver-rules=MAP chat.hl.internal 10.0.2.101,MAP sheets.hl.internal 10.0.2.102'"
            in (launch_line)
        )
        # The start page is the first (active) tab, then one tab per app, in order.
        expected = "file:///home/ga/Desktop/APPS.html 'http://chat.hl.internal/?sid=1' 'http://sheets.hl.internal/?sid=1' "
        assert expected in launch_line
        assert "127.0.0.1:12345:127.0.0.1:9222" in popen_calls[0]
    else:
        with pytest.raises(RuntimeError, match="No X desktop"):
            vm.open_browser(tabs, hosts)
        assert not popen_calls


def test_browser_without_app_names_still_opens_the_start_page(monkeypatch):
    vm = object.__new__(hub_vm.HubWorkerVM)
    vm.worker, vm.cdp_port = "boss", 12345
    vm.ssh_base, vm.env = ["ssh", "ga@127.0.0.1"], {}
    commands = []

    def run(command):
        commands.append(command)
        return SimpleNamespace(stdout="X0" if command == "ls /tmp/.X11-unix" else "")

    vm.run = run
    monkeypatch.setattr(hub_vm.HubWorkerVM, "connect_browser", lambda self: self.cdp_port)
    assert vm.open_browser() == 12345
    assert "--host-resolver-rules" not in commands[-1]
    assert "--user-data-dir=/home/ga/worker-browser file:///home/ga/Desktop/APPS.html >" in commands[-1]


def test_reset_retains_overlays_when_stop_fails(company, monkeypatch):
    launch(company)
    overlay = company / "runtime/vms/boss/worker.qcow2"
    overlay.write_text("owned disk")

    def fail(_):
        raise OSError("cannot stop")

    monkeypatch.setattr(hub_vm, "_stop_process", fail)
    with pytest.raises(RuntimeError, match="cannot stop"):
        hub_vm.reset_company(company)
    assert overlay.read_text() == "owned disk"


@pytest.mark.parametrize("kind", ["copy", "hardlink", "parent_symlink"])
def test_private_task_materials_and_parent_symlinks_rejected(company, kind):
    import shutil

    root = company / "world/materials"
    private = company / "tasks/task/workflow.json"
    if kind == "copy":
        shutil.copyfile(private, root / "analyst/workflow.json")
    elif kind == "hardlink":
        (root / "analyst/innocent.txt").hardlink_to(private)
    else:
        moved = company / "private-materials"
        root.rename(moved)
        root.symlink_to(moved, target_is_directory=True)
    with pytest.raises(ValueError, match="private|regular"):
        launch(company)
    assert not (company / "runtime/vms").exists()


@pytest.mark.parametrize(
    "bad", ["http://localhost:8000/go", "http://192.168.1.20:8100", "http://192.168.1.20:8100/?other=1"]
)
def test_material_url_checks_compare_origins_not_literal_queries(company, bad):
    (company / "world/materials/analyst/notes/onboarding.md").write_text(bad)
    with pytest.raises(ValueError, match="private/unassigned"):
        launch(company)


def test_vm_network_restricts_egress_to_assigned_proxies(company, fake_vm):
    real_branch(company)
    for vm in fake_vm.instances:
        network = vm.command[vm.command.index("-netdev") + 1]
        assert "restrict=on" in network
        assert ":8000" not in network and ":8001" not in network
        assert "10.0.2.100" not in network
        if vm.worker == "analyst":
            assert ":8100" not in network and ":8200" not in network and ":8201" not in network
            assert "guestfwd=tcp:10.0.2.101:80-cmd:" in network and "127.0.0.1 8101" in network
            assert "10.0.2.102" not in network
        else:
            assert "guestfwd=tcp:10.0.2.101:80-cmd:" in network and "127.0.0.1 8100" in network
            assert "guestfwd=tcp:10.0.2.102:80-cmd:" in network and "127.0.0.1 8200" in network


def test_a_document_keeps_every_line_under_its_heading():
    """A heading used to swallow its block, so the artifact lost what identified the document.

    An invoice rendered as the supplier's name alone: the number, both dates and the customer were
    dropped from the PDF the worker opens. 151 occurrences across 130 of 1,027 documents, and no
    gate saw it because every gate reads the source text rather than the rendered file.
    """
    from company_envs.world.documents import _paragraphs

    invoice = (
        "# Sandstone Printer Care\n"
        "Invoice 4471\n"
        "Issued August 12, 2026\n"
        "Customer: Open Horizon Motorhome Rentals\n\n"
        "Service and parts: $296.00\nCall-out: $95.00\n"
    )
    blocks = list(_paragraphs(invoice))
    assert blocks[0] == ("h", 1, "Sandstone Printer Care")
    rendered = " ".join(text for _kind, _level, text in blocks)
    for kept in ("Invoice 4471", "August 12, 2026", "Open Horizon Motorhome Rentals", "$296.00"):
        assert kept in rendered

    # Every line keeps its own paragraph, so a label and its value never run together.
    assert ("p", 0, "Call-out: $95.00") in blocks
    assert ("p", 0, "Customer: Open Horizon Motorhome Rentals") in blocks
    # A paragraph that is one long line is untouched: the sources do not hard-wrap.
    prose = "One long sentence about the review that is not wrapped in the source at all."
    assert list(_paragraphs(prose)) == [("p", 0, prose)]


def test_materials_render_into_the_formats_their_names_promise():
    from company_envs.world.documents import render_material

    pdf = render_material(
        "Downloads/Invoice 4471.pdf",
        b"# Invoice 4471\n\nHarbor Supply, due September 30, 2026.\n\nTotal $4,200.",
    )
    assert pdf[:4] == b"%PDF"
    docx = render_material("Documents/Policy.docx", b"# Rate policy\n\nEffective July 1.")
    assert docx[:2] == b"PK"
    xlsx = render_material(
        "Documents/Q3 forecast v2.xlsx",
        b"## Sheet: Forecast\nmonth,amount\nJuly,4200\nAugust,3900.5\n## Sheet: Notes\nnote\nprovisional\n",
    )
    assert xlsx[:2] == b"PK"
    import io

    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(xlsx))
    assert book.sheetnames == ["Forecast", "Notes"] and book["Forecast"]["B2"].value == 4200
    pptx = render_material(
        "Desktop/Board update.pptx", b"# Q3 update\n- revenue up\n---\n# Risks\n- one supplier"
    )
    assert pptx[:2] == b"PK"
    assert render_material("notes/todo.md", b"- call Jo") == b"- call Jo"
    assert render_material("a.pdf", b"%PDF-1.4 already real") == b"%PDF-1.4 already real"


@pytest.mark.parametrize("stopped", [False, True])
def test_snapshot_reset_reuses_disks_and_skips_setup(company, fake_vm, monkeypatch, stopped):
    first = real_branch(company)
    root = company / "runtime/vms"
    monkeypatch.setattr(hub_vm, "_stop_process", lambda _: None)
    monkeypatch.setattr(hub_vm, "_has_snapshot", lambda _: True)
    if stopped:
        hub_vm.stop_company(company)
    write(root / "boss/trace/receipt.json", {"old": "episode"})
    overlay = root / "boss/worker.qcow2"
    inode = overlay.stat().st_ino
    before = fake_vm.instances.copy()
    network = before[0].command.copy()
    # Repeating reset must keep the original snapshot, rather than save episode work.
    for _ in range(2):
        report = hub_vm.reset_company(company)
        assert report["status"] == "running" and report["reset_mode"] == "snapshot"
        assert report["boot_seconds"] == first["boot_seconds"]
        assert report["reset_seconds"] > 0
        assert report == read(root / "LAUNCH.json")
        assert overlay.stat().st_ino == inode
        assert not (root / "boss/trace").exists()
        assert all(vm.restored and not vm.saved for vm in fake_vm.instances[-2:])
        assert fake_vm.instances[-2].command == network
        assert read(root / "boss/vm.json")["vm"]["tunnel"]["pid"] == 789
    # Fixture measurement: two fresh boots, four uploads and two browser launches
    # become two memory restores with no uploads or browser launches on reset.
    after = fake_vm.instances[-2:]
    assert sum(not vm.restored for vm in before) == 2
    assert sum(not vm.restored for vm in after) == 0
    assert sum(len(vm.uploads) for vm in before) == 4
    assert sum(len(vm.uploads) for vm in after) == 0
    assert sum(vm.opened for vm in before) == 2
    assert sum(vm.opened for vm in after) == 0


@pytest.mark.parametrize(
    "change", ["legacy", "missing", "receipt", "disk", "materials", "brief", "url", "clock"]
)
def test_snapshot_reset_falls_back_for_missing_or_stale_setup(company, fake_vm, monkeypatch, change):
    real_branch(company)
    root = company / "runtime/vms"
    monkeypatch.setattr(hub_vm, "_stop_process", lambda _: None)
    monkeypatch.setattr(hub_vm, "_has_snapshot", lambda _: change != "disk")
    if change == "legacy":
        report = read(root / "LAUNCH.json")
        del report["snapshot_key"]
        write(root / "LAUNCH.json", report)
    elif change == "missing":
        receipt = read(root / "boss/vm.json")
        del receipt["snapshot"]
        write(root / "boss/vm.json", receipt)
    elif change == "receipt":
        hub_vm.stop_company(company)
        (root / "boss/vm.json").unlink()
    elif change == "materials":
        (company / "world/materials/boss/notes/onboarding.md").write_text("Updated onboarding")
    elif change == "brief":
        write(company / "tasks/task/assignment.json", {"brief": "Check the revised deliveries."})
    elif change == "url":
        endpoints = read(company / "runtime/endpoints.json")
        endpoints["apps"]["chat"]["workers"]["boss"] = "http://localhost:9100/?sid=next-episode"
        write(company / "runtime/endpoints.json", endpoints)
    elif change == "clock":
        write(company / "world/SEED.json", {"reference_date": "2026-09-09"})
    report = hub_vm.reset_company(company)
    assert report["reset_mode"] == "fresh_boot"
    assert report["reset_seconds"] > 0
    assert all(not vm.restored and vm.saved for vm in fake_vm.instances[-2:])
    if change == "materials":
        assert (root / "boss/guest/Desktop/notes/onboarding.md").read_text() == "Updated onboarding"
    if change == "url":
        assert "127.0.0.1 9100" in fake_vm.instances[-2].command[2]


@pytest.mark.parametrize("failure", ["boot", "browser"])
def test_snapshot_restore_failure_stops_every_owned_vm(company, fake_vm, monkeypatch, failure):
    real_branch(company)
    monkeypatch.setattr(hub_vm, "_stop_process", lambda _: None)
    monkeypatch.setattr(hub_vm, "_has_snapshot", lambda _: True)
    if failure == "boot":
        fake_vm.fail = True
    else:

        def connect(vm):
            if vm.worker == "analyst":
                raise RuntimeError("browser failed")

        monkeypatch.setattr(fake_vm, "connect_browser", connect)
    with pytest.raises(RuntimeError, match=failure):
        hub_vm.reset_company(company)
    report = read(company / "runtime/vms/LAUNCH.json")
    assert report["status"] == "failed"
    assert all(vm.closed for vm in fake_vm.instances[-2:])
    assert all((vm.directory / "worker.qcow2").is_file() for vm in fake_vm.instances[-2:])


def test_snapshot_save_failure_fails_launch_and_cleans_up(company, fake_vm, monkeypatch):
    def save(vm):
        raise RuntimeError("snapshot failed")

    monkeypatch.setattr(fake_vm, "save_snapshot", save)
    with pytest.raises(RuntimeError, match="snapshot failed"):
        real_branch(company)
    assert all(vm.closed for vm in fake_vm.instances)
    root = company / "runtime/vms"
    assert read(root / "LAUNCH.json")["status"] == "failed"
    assert "snapshot" not in read(root / "boss/vm.json")


@pytest.mark.parametrize(
    "snapshots, expected",
    [
        ([], False),
        ([{"name": "another", "vm-state-size": 1024}], False),
        ([{"name": "episode-start", "vm-state-size": 0}], False),
        ([{"name": "episode-start", "vm-state-size": 1024}], True),
    ],
)
def test_snapshot_requires_saved_memory(tmp_path, monkeypatch, snapshots, expected):
    import json

    (tmp_path / "worker.qcow2").touch()

    def run(command, **kwargs):
        assert command == ["qemu-img", "info", "--output=json", str(tmp_path / "worker.qcow2")]
        assert kwargs["check"] and kwargs["timeout"] == 30
        return SimpleNamespace(stdout=json.dumps({"snapshots": snapshots}))

    monkeypatch.setattr(hub_vm.subprocess, "run", run)
    assert hub_vm._has_snapshot(tmp_path) is expected
    (tmp_path / "worker.qcow2").unlink()
    assert not hub_vm._has_snapshot(tmp_path)


def test_reset_uses_source_materials_without_rendering_again(company, fake_vm, monkeypatch):
    material = company / "world/materials/boss/forecast.xlsx"
    material.write_text("month,amount\nJuly,4200\n")
    real_branch(company)
    delivered = company / "runtime/vms/boss/guest/Desktop/forecast.xlsx"
    assert delivered.read_bytes().startswith(b"PK")
    monkeypatch.setattr(hub_vm, "_stop_process", lambda _: None)
    monkeypatch.setattr(hub_vm, "_has_snapshot", lambda _: True)

    def render(*args):
        pytest.fail("Unchanged source materials must reuse the saved desktop without rendering")

    monkeypatch.setattr(hub_vm, "render_material", render)
    assert hub_vm.reset_company(company)["reset_mode"] == "snapshot"


def test_launch_and_reset_record_elapsed_time(company, fake_vm, monkeypatch):
    elapsed = 0.0

    def advance(seconds):
        nonlocal elapsed
        elapsed += seconds

    def timed_method(name, seconds, restore_seconds=None):
        original = getattr(fake_vm, name)

        def call(*args, **kwargs):
            result = original(*args, **kwargs)
            advance(restore_seconds if kwargs.get("snapshot") else seconds)
            return result

        monkeypatch.setattr(fake_vm, name, call)

    monkeypatch.setattr(hub_vm.time, "monotonic", lambda: elapsed)
    monkeypatch.setattr(hub_vm, "_stop_process", lambda _: advance(0.25))
    monkeypatch.setattr(hub_vm, "_has_snapshot", lambda _: True)
    timed_method("boot", 40, 3)
    timed_method("upload", 2)
    timed_method("open_browser", 5)
    timed_method("connect_browser", 1)
    timed_method("save_snapshot", 4)
    report = real_branch(company)
    assert report["boot_seconds"] == 98  # Two desktops: boot, two uploads, then browser readiness.
    assert report["snapshot_seconds"] == 8
    assert report["reset_seconds"] is None
    report = hub_vm.reset_company(company)
    assert report["boot_seconds"] == 98
    assert report["reset_seconds"] == 9  # Stop four host processes, then restore both desktops.
    assert report == read(company / "runtime/vms/LAUNCH.json")
    assert read(company / "runtime/vms/boss/vm.json")["reset_seconds"] == 4
    # A missing snapshot times the entire fresh launch, including saving its replacement.
    monkeypatch.setattr(hub_vm, "_has_snapshot", lambda _: False)
    report = hub_vm.reset_company(company)
    assert report["reset_mode"] == "fresh_boot"
    assert report["reset_seconds"] == 107


def test_restore_cleanup_failure_preserves_error_and_attempts_every_vm(company, fake_vm, monkeypatch):
    real_branch(company)
    monkeypatch.setattr(hub_vm, "_stop_process", lambda _: None)
    monkeypatch.setattr(hub_vm, "_has_snapshot", lambda _: True)
    fake_vm.fail = True

    def close(vm):
        vm.closed = True
        if vm.worker == "boss":
            raise OSError("cleanup failed")

    monkeypatch.setattr(fake_vm, "close", close)
    with pytest.raises(RuntimeError, match="boot failed"):
        hub_vm.reset_company(company)
    assert all(vm.closed for vm in fake_vm.instances[-2:])
    assert read(company / "runtime/vms/LAUNCH.json")["cleanup_errors"] == ["boss: cleanup failed"]


def test_invalid_handoff_leaves_saved_vms_untouched(company, fake_vm, monkeypatch):
    first = real_branch(company)

    def stop(_):
        pytest.fail("Invalid handoffs must fail before stopping the current desktop")

    monkeypatch.setattr(hub_vm, "_stop_process", stop)
    (company / "runtime/endpoints.json").unlink()
    with pytest.raises(FileNotFoundError, match="hub-serve"):
        hub_vm.reset_company(company)
    assert read(company / "runtime/vms/LAUNCH.json") == first


def test_snapshot_inspection_failure_keeps_the_overlay(company, fake_vm, monkeypatch):
    real_branch(company)
    root = company / "runtime/vms"
    monkeypatch.setattr(hub_vm, "_stop_process", lambda _: None)

    def inspect(_):
        raise OSError("cannot read snapshot")

    monkeypatch.setattr(hub_vm, "_has_snapshot", inspect)
    with pytest.raises(OSError, match="cannot read snapshot"):
        hub_vm.reset_company(company)
    assert (root / "boss/worker.qcow2").read_text() == "overlay"
    assert len(fake_vm.instances) == 2
    assert read(root / "LAUNCH.json")["status"] == "stopped"


def test_a_stale_handoff_is_named_before_any_guest_boots(company, monkeypatch):
    """runtime/endpoints.json names this serving process's ports, so one left behind is dead.

    The guest cannot tell: QEMU's guestfwd accepts the connection before its host-side ``nc`` runs,
    so a dead port comes back as a connection *reset* and reads as a broken app. One live VM run was
    spent on a handoff an earlier ``hub-serve --check`` had left behind. Checked from the host, where
    the answer is unambiguous, and before a single overlay is created.
    """
    monkeypatch.setattr(hub_vm.shutil, "which", lambda binary: "/usr/bin/" + binary)
    monkeypatch.setattr(hub_vm.os, "access", lambda *args: True)
    monkeypatch.setattr(hub_vm, "_files", lambda root: [])
    (company / "browser").mkdir()
    (company / "browser/chrome").write_text("browser double")
    (company / "base.qcow2").write_text("image double")
    with socket.socket() as spare:
        spare.bind(("127.0.0.1", 0))
        dead = spare.getsockname()[1]
    plans = {"boss": {"addresses": {"chat": {"proxy_port": dead}}}}
    with pytest.raises(ConnectionRefusedError, match="stale hub-serve handoff"):
        hub_vm._preflight(company / "base.qcow2", company / "browser", company / "runtime/vms", plans)


def test_live_proxy_ports_reports_only_what_nothing_is_serving():
    with socket.socket() as listening:
        listening.bind(("127.0.0.1", 0))
        listening.listen(1)
        live = listening.getsockname()[1]
        with socket.socket() as spare:
            spare.bind(("127.0.0.1", 0))
            dead = spare.getsockname()[1]
        plans = {
            "boss": {"addresses": {"chat": {"proxy_port": live}, "sheets": {"proxy_port": dead}}},
            "analyst": {"addresses": {"chat": {"proxy_port": live}, "notes": {}}},
        }
        assert hub_vm.live_proxy_ports(plans) == [f"boss/sheets:{dead}", "analyst/notes: no proxy port"]


def test_app_readiness_retries_a_reset_and_allows_a_slow_first_response():
    """``--retry-connrefused`` never covered the failure the guest actually sees, and a three-second
    cap meant an app whose first response took longer could never be declared ready however many
    times it retried. Measured: one guest reads a 6.0 MB state in 2.2 s and three concurrent guests
    take 20-23 s for the same read, with 51 proxies and 17 vite servers in one host process at 89%
    CPU. The loop is shell so it does not depend on the guest curl supporting --retry-all-errors.
    """
    command = hub_vm.app_ready_command("http://chat.hl.internal/?sid=1", "chat.hl.internal", "10.0.2.101")
    assert f"--max-time {hub_vm.READY_MAX_SECONDS}" in command
    assert hub_vm.READY_MAX_SECONDS > 3
    assert f"seq 1 {hub_vm.READY_ATTEMPTS - 1}" in command
    assert f"sleep {hub_vm.READY_DELAY}" in command
    assert "--resolve chat.hl.internal:80:10.0.2.101" in command
    # The last attempt runs outside the loop, so its message is the one the launcher raises.
    assert command.count("curl") == 2 and command.count("exit 0") == 1
    assert command.rstrip().endswith(">/dev/null")


def test_the_guest_settle_comes_from_the_render_report(company):
    """A guest needs roughly three times the host's hydration time: bamboohr shows 3,172 characters
    and no seeded strings four seconds after load in a guest and 198,993 -- the host's own figure --
    after twelve. The render gate measures it on the host and publishes the guest figure, because its
    own browser never runs in a guest.
    """
    assert hub_vm.guest_settle(company) == float(hub_vm.GUEST_SETTLE_SECONDS)
    assert hub_vm.GUEST_SETTLE_SECONDS >= 12  # the guest's measured figure, not the host's
    write(company / "runtime/RENDER.json", {"ok": True, "guest_settle_seconds": 9})
    assert hub_vm.guest_settle(company) == 9.0
    write(company / "runtime/RENDER.json", {"ok": True, "guest_settle_seconds": 10_000})
    assert hub_vm.guest_settle(company) == float(hub_vm.GUEST_SETTLE_SECONDS)
    write(company / "runtime/RENDER.json", {"ok": True, "guest_settle_seconds": True})
    assert hub_vm.guest_settle(company) == float(hub_vm.GUEST_SETTLE_SECONDS)


def test_a_worker_is_not_ready_until_its_pages_have_hydrated(company, fake_vm):
    """CDP answering and the start page appearing say the browser is up, not that the apps hold the
    world. Waited on the guest's own clock, and before any snapshot, because a snapshot taken of
    unhydrated apps freezes them that way for every reset.
    """
    write(company / "runtime/RENDER.json", {"ok": True, "guest_settle_seconds": 7})
    real_branch(company)
    vm = fake_vm.instances[0]
    assert "sleep 7" in vm.commands
    assert vm.commands.index("sleep 7") > max(
        index for index, command in enumerate(vm.commands) if "curl" in command
    )
    assert read(company / "runtime/vms/boss/vm.json")["settle_seconds"] == 7.0
