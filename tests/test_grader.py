"""Offline rewards using small app fixtures and read-only doubles."""

import json
from copy import deepcopy

import pytest
from pydantic import ValidationError

from company_envs.storage import digest, read, write
from company_envs.world.grader import (
    JUDGE_VOTES,
    AppPath,
    CalibrationError,
    Check,
    Judgment,
    Predicate,
    TaskGrader,
    TaskJudgment,
    author_grader,
    authoring_payload,
    calibrate,
    evaluate_predicate,
    grade,
    select,
)

APP = "Zendesk_mock"
TASK = "resolve"
SID = "episode-1"


class Client:
    def __init__(self, initial):
        self.initial = deepcopy(initial)
        self.current = deepcopy(initial)

    def inspect(self, sid):
        assert sid == SID
        return {"initial_state": self.initial, "current_state": self.current, "state_diff": {}}


class Model:
    def __init__(self, result):
        self.result = result
        self.calls = []

    def call(self, job, prompt, response_type):
        self.calls.append((job, prompt, response_type))
        assert isinstance(self.result, response_type)
        result = self.result
        if job == "task_judgment":
            # The grader asks for one check per call; answer only the check it named.
            packet = json.loads(prompt[prompt.index("{") :])
            wanted = {c["check"]["id"] for c in packet["checks"]}
            result = TaskJudgment(checks=[c for c in result.checks if c.check_id in wanted])
        return result, {"job": job, "session_id": f"fresh-{job}"}


def predicate(selector, operator="equals", value="solved"):
    return Predicate(
        app_id=APP,
        selector=selector,
        operator=operator,
        value_json=json.dumps(value) if operator in {"equals", "contains", "count_gte"} else None,
    )


def state_check(index, selector, operator="equals", value="solved", **kwargs):
    return Check(
        id=f"c{index}",
        criterion_ref=f"/success_criteria/{index}",
        kind="state",
        description="Verify required business change",
        predicate=predicate(selector, operator, value),
        **kwargs,
    )


@pytest.fixture
def initial():
    return {
        "tickets": [
            {
                "id": 1310,
                "status": "open",
                "tags": ["pickup"],
                "subject": "Pickup H310",
                "satisfaction_rating": None,
                "is_public": True,
            },
            {"id": 1311, "status": "open"},
            {"id": 1312, "status": "open"},
        ],
        "users": [{"id": 1, "name": "Case owner"}],
        "currentUser": {"id": 1},
        "comments": {"1310": []},
        "ui": {},
    }


@pytest.fixture
def company(tmp_path, initial):
    folder = tmp_path / "companies" / "demo"
    write(folder / "world" / f"{APP}.state.json", initial)
    write(
        folder / "apps.json",
        {
            "apps": [
                {
                    "app_id": APP,
                    "top_level_keys": list(initial),
                    "schema": "schema.md",
                    "state_file": f"world/{APP}.state.json",
                }
            ]
        },
    )
    (tmp_path / "schema.md").write_text("Zendesk: tickets have id, status, tags.")
    write(folder / "runtime" / "sessions.json", {"sid": SID})
    write(folder / "tasks" / TASK / "assignment.json", {"brief": "Resolve the target case."})
    write(
        folder / "tasks" / TASK / "workflow.json",
        {
            "success_criteria": [{"requirement": "Resolve target", "method": "state"}],
            "worker_ids": ["w1"],
            "completion": {"feasible_path": {"secret": "HIDDEN-SOLUTION"}, "outcomes": [{"id": "done"}]},
        },
    )
    return tmp_path, folder, {APP: Client(initial)}


def prepare(company, checks=None, workers=None):
    root, folder, clients = company
    checks = checks or [state_check(0, "$.tickets[?(@.id==1310)].status")]
    workflow = read(folder / "tasks" / TASK / "workflow.json")
    workflow["success_criteria"] = [{"requirement": c.description, "method": c.kind} for c in checks]
    if workers is not None:
        workflow["worker_ids"] = workers
    write(folder / "tasks" / TASK / "workflow.json", workflow)
    model = Model(TaskGrader(checks=checks))
    author_grader(root, folder, TASK, models=model)
    return folder, clients, model


def golden(clients):
    result = deepcopy(clients[APP].initial)
    result["tickets"][0]["status"] = "solved"
    return result


def certify(folder, clients, state=None):
    state = state or golden(clients)
    return calibrate(folder, TASK, clients, reference_diff={APP: {"tickets": state["tickets"]}})


@pytest.mark.parametrize(
    ("selector", "operator", "value", "expected"),
    [
        ("$.tickets[?(@.id==1310)].status", "equals", "open", True),
        ("$.tickets[?(@.id==1310)].status", "equals", "solved", False),
        ("$.tickets[?(@.id==1310)].tags", "contains", "pickup", True),
        ("$.tickets[?(@.id==1310)].subject", "contains", "H310", True),
        ("$.tickets", "count_gte", 3, True),
        ("$.tickets[*].id", "count_gte", 3, True),
        ("$.tickets[?(@.id==1310)]", "count_gte", 2, False),
        ("$.comments", "contains", "1310", True),
        ('$.comments["1310"]', "exists", None, True),
        ("$.tickets[0].satisfaction_rating", "equals", None, True),
        ("$.tickets[0].satisfaction_rating", "exists", None, True),
        ("$.tickets[0].missing", "equals", None, False),
        ("$.tickets[0].missing", "unchanged", None, False),
        ("$.tickets[?(@.id==99999)].status", "exists", None, False),
        ("$.tickets[*].status", "equals", "open", False),
        ("$.tickets[0].is_public", "equals", 1, False),
    ],
)
def test_predicates(initial, selector, operator, value, expected):
    assert (
        evaluate_predicate(predicate(selector, operator, value), {APP: initial}, {APP: initial}) is expected
    )


def test_changed_deletion_and_unchanged(initial):
    final = deepcopy(initial)
    final["tickets"][0]["status"] = "solved"
    assert evaluate_predicate(predicate("$.tickets[0].status", "changed"), {APP: initial}, {APP: final})
    assert evaluate_predicate(predicate("$.users", "unchanged"), {APP: initial}, {APP: final})
    del final["tickets"][0]["status"]
    assert evaluate_predicate(predicate("$.tickets[0].status", "changed"), {APP: initial}, {APP: final})
    assert not evaluate_predicate(predicate("$.tickets[0].status", "unchanged"), {APP: initial}, {APP: final})


@pytest.mark.parametrize(
    "selector", ["tickets", "$..tickets", "$.tickets[-1]", "$.tickets[?(__import__('os'))]", "$"]
)
def test_invalid_selectors(selector):
    with pytest.raises(ValueError):
        predicate(selector)


def test_noop_calibration_rejects_any_passing_check(company):
    folder, clients, _ = prepare(
        company, [state_check(0, "$.tickets[?(@.id==1310)].status"), state_check(1, "$.users", "exists")]
    )
    with pytest.raises(CalibrationError, match="EVERY"):
        certify(folder, clients)
    proof = read(folder / "runtime" / "grades" / TASK / "calibration.json")
    assert not proof["accepted"] and proof["initial_score"] == 0.5 and proof["reference_score"] == 1
    with pytest.raises(CalibrationError):
        grade(folder, TASK, clients)


def test_golden_diff_and_noop_grade(company):
    folder, clients, _ = prepare(company)
    proof = certify(folder, clients)
    assert proof["initial_score"] == 0 and proof["reference_score"] == 1
    assert clients[APP].current == clients[APP].initial
    assert grade(folder, TASK, clients)["score"] == 0
    clients[APP].current = golden(clients)
    report = grade(folder, TASK, clients)
    assert report["score"] == 1 and report["passed"]
    assert read(folder / "runtime" / "grades" / TASK / "report.json")["score"] == 1


def test_partial_wrong_target_and_guard_failure(company):
    folder, clients, _ = prepare(
        company,
        [
            state_check(0, "$.tickets[?(@.id==1310)].status", guards=[predicate("$.users", "unchanged")]),
            state_check(1, "$.tickets[?(@.id==1311)].status"),
        ],
    )
    final = golden(clients)
    final["tickets"][1]["status"] = "solved"
    certify(folder, clients, final)
    clients[APP].current["tickets"][2]["status"] = "solved"
    assert grade(folder, TASK, clients)["score"] == 0
    clients[APP].current = golden(clients)
    assert grade(folder, TASK, clients)["score"] == 0.5
    clients[APP].current["users"] = []
    assert grade(folder, TASK, clients)["score"] == 0


def test_authoring_information_barrier(company):
    root, folder, _ = company
    payload = authoring_payload(root, folder, TASK)
    assert set(payload) == {"public_brief", "apps", "success_criteria", "contract"}
    assert payload["apps"][0]["initial_state"]["tickets"][0]["status"] == "open"
    _, _, model = prepare(company)
    job, prompt, response_type = model.calls[0]
    assert job == "task_grader" and response_type is TaskGrader
    data = json.loads(prompt.split("\n")[-1])
    assert "feasible_path" not in json.dumps(data) and "HIDDEN-SOLUTION" not in prompt
    assert "completion" not in data and "receipt" not in data


def test_requires_calibration_and_pins_baseline_and_grader(company):
    folder, clients, _ = prepare(company)
    with pytest.raises(CalibrationError, match="calibrate"):
        grade(folder, TASK, clients)
    certify(folder, clients)
    clients[APP].initial["ui"]["extra"] = True
    with pytest.raises(CalibrationError, match="baseline"):
        grade(folder, TASK, clients)
    del clients[APP].initial["ui"]["extra"]
    path = folder / "tasks" / TASK / "grader.json"
    draft = read(path)
    draft["grader"]["checks"][0]["description"] = "Edited grader"
    write(path, draft)
    with pytest.raises(CalibrationError, match="stale"):
        grade(folder, TASK, clients)


def test_reference_required_and_invalid_reference_rejected(company):
    folder, clients, _ = prepare(company)
    clients[APP].current = golden(clients)
    with pytest.raises(FileNotFoundError):
        calibrate(folder, TASK, clients)
    with pytest.raises(CalibrationError):
        calibrate(folder, TASK, clients, reference_states={APP: clients[APP].initial})
    write(folder / "tasks" / TASK / "reference.json", {"final_states": {APP: golden(clients)}})
    assert calibrate(folder, TASK, clients)["accepted"]


def test_three_worker_contributions_filter_sessions_and_irrelevant_keys(company):
    folder, clients, _ = prepare(company, workers=["w1", "w2", "w3"])
    certify(folder, clients)
    clients[APP].current = golden(clients)
    path = folder / "runtime" / "attribution" / f"{APP}.jsonl"
    path.parent.mkdir(parents=True)
    rows = [
        {"sid": SID, "worker_id": "w1", "changed_keys": ["tickets"], "at": "2026-09-07T15:00:00Z"},
        {"sid": SID, "worker_id": "w2", "changed_keys": ["ui"], "at": "2026-09-07T15:01:00Z"},
        {"sid": "old-episode", "worker_id": "w3", "changed_keys": ["tickets"], "at": "old"},
        {"sid": SID, "worker_id": "outsider", "changed_keys": ["tickets"], "at": "now"},
    ]
    path.write_text("\n".join(map(json.dumps, rows)) + "\ninvalid-json\n")
    report = grade(folder, TASK, clients)
    assert report["score"] == pytest.approx(1 / 3)
    assert not report["passed"] and report["contributions"]["observed_workers"] == 1
    assert report["contributions"]["warnings"]
    rows.extend(
        {"sid": SID, "worker_id": worker, "changed_keys": ["tickets"], "at": "later"}
        for worker in ("w2", "w3")
    )
    path.write_text("\n".join(map(json.dumps, rows)))
    report = grade(folder, TASK, clients)
    assert report["score"] == 1 and report["passed"]
    assert all(row["passing_checks"] == ["c0"] for row in report["contributions"]["workers"])
    assert "collection-level" in report["contributions"]["basis"]


def semantic_check():
    return Check(
        id="c1",
        criterion_ref="/success_criteria/1",
        kind="artifact",
        description="Explain supported resolution",
        rubric="Require supported final terms.",
        app_paths=[AppPath(app_id=APP, selector="$.tickets[?(@.id==1310)]")],
        material_paths=["w1/resolution.txt"],
    )


def test_artifact_content_packet_and_separate_judgment(company):
    _, folder, _ = company
    material = folder / "world" / "materials" / "w1" / "resolution.txt"
    material.parent.mkdir(parents=True)
    material.write_text("Initial draft")
    folder, clients, author = prepare(company, [state_check(0, "$.tickets[0].status"), semantic_check()])
    certify(folder, clients)
    assert grade(folder, TASK, clients)["score"] == 0
    clients[APP].current = golden(clients)
    content = "Final supported terms\nLiteral \\n stays literal"
    material.write_text(content)
    report = grade(folder, TASK, clients)
    assert report["score"] == 0.6 and report["checks"][1]["status"] == "pending"
    judge = Model(
        TaskJudgment(
            checks=[
                Judgment(
                    check_id="c1",
                    status="pass",
                    reason="Supported terms",
                    evidence_refs=["app:0", "material:0"],
                )
            ]
        )
    )
    report = grade(folder, TASK, clients, models=judge)
    assert report["score"] == 1 and report["passed"]
    assert len(author.calls) == 1 and judge.calls[0][0] == "task_judgment"
    packet = read(folder / "runtime" / "grades" / TASK / "evidence.json")
    assert packet["checks"][0]["evidence"][1]["content"] == content
    assert "HIDDEN-SOLUTION" not in judge.calls[0][1]
    material.unlink()
    report = grade(folder, TASK, clients, models=judge)
    assert report["score"] == 0.6 and report["checks"][1]["status"] == "fail"
    assert len(judge.calls) == JUDGE_VOTES  # one bounded call per vote, all for the single pending check


def test_untrusted_judge_and_missing_hub_fail_closed(company):
    check = semantic_check().model_copy(update={"material_paths": []})
    folder, clients, _ = prepare(company, [state_check(0, "$.tickets[0].status"), check])
    certify(folder, clients)
    clients[APP].current = golden(clients)
    judge = Model(
        TaskJudgment(
            checks=[
                Judgment(check_id="c1", status="pass", reason="Claims completion", evidence_refs=["invented"])
            ]
        )
    )
    report = grade(folder, TASK, clients, models=judge)
    assert report["score"] == 0.6 and report["checks"][1]["status"] == "pending"
    assert any(key.startswith("judgment") for key in report["errors"])
    report = grade(folder, TASK, {})
    assert report["score"] == 0 and APP in report["errors"]


def test_invalid_contracts(company):
    with pytest.raises(ValidationError):
        Check(**(semantic_check().model_dump() | {"material_paths": ["../../workflow.json"]}))
    with pytest.raises(ValidationError):
        TaskGrader(checks=[state_check(0, "$.tickets"), state_check(0, "$.tickets")])
    with pytest.raises(ValidationError):
        predicate("$.tickets", "count_gte", 0)
    with pytest.raises(ValidationError):
        predicate("$.tickets", "equals", float("nan"))
    root, folder, _ = company
    with pytest.raises(ValueError, match="cover every"):
        author_grader(root, folder, TASK, models=Model(TaskGrader(checks=[state_check(1, "$.tickets")])))


def test_selector_exact_types_and_escaped_keys():
    state = {"rows": [{"id": True}, {"id": 1}], "a.b": {"line": "value"}}
    assert select(state, "$.rows[?(@.id==true)]") == [{"id": True}]
    assert select(state, "$.rows[?(@.id==1)]") == [{"id": 1}]
    assert select(state, '$["a.b"].line') == ["value"]


def test_symlink_material_and_task_path_escape(company, tmp_path):
    _, folder, _ = company
    check = semantic_check()
    folder, clients, _ = prepare(company, [state_check(0, "$.tickets[0].status"), check])
    certify(folder, clients)
    clients[APP].current = golden(clients)
    material = folder / "world" / "materials" / "w1" / "resolution.txt"
    material.parent.mkdir(parents=True)
    secret = tmp_path / "private.txt"
    secret.write_text("SHOULD-NOT-BE-READ")
    material.symlink_to(secret)
    report = grade(folder, TASK, clients)
    assert report["checks"][1]["status"] == "fail"
    packet = read(folder / "runtime" / "grades" / TASK / "evidence.json")
    assert "SHOULD-NOT-BE-READ" not in json.dumps(packet)
    with pytest.raises(ValueError, match="invalid task"):
        grade(folder, "../private", clients)


def test_broken_attribution_cannot_earn_credit(company):
    folder, clients, _ = prepare(company, workers=["w1", "w2", "w3"])
    certify(folder, clients)
    clients[APP].current = golden(clients)
    path = folder / "runtime" / "attribution" / f"{APP}.jsonl"
    path.parent.mkdir(parents=True)
    path.write_text(json.dumps({"sid": SID, "worker_id": "w1", "changed_keys": ["tickets"]}))
    report = grade(folder, TASK, clients)
    assert report["contributions"]["observed_workers"] == 0
    assert report["contributions"]["warnings"]


def test_unchanged_semantic_evidence_cannot_be_judged_pass(company):
    check = Check(
        id="c1",
        criterion_ref="/success_criteria/1",
        kind="judgment",
        description="Assess an updated report",
        rubric="A complete report is required.",
        app_paths=[AppPath(app_id=APP, selector="$.comments[*]")],
    )
    folder, clients, _ = prepare(company, [state_check(0, "$.tickets[0].status"), check])
    certify(folder, clients)
    clients[APP].current = golden(clients)
    judge = Model(
        TaskJudgment(
            checks=[Judgment(check_id="c1", status="pass", reason="Unsupported", evidence_refs=["app:0"])]
        )
    )
    report = grade(folder, TASK, clients, models=judge)
    assert report["score"] == 0.6 and report["checks"][1]["status"] == "fail"
    assert not judge.calls


def test_default_author_model_configuration_and_schema(company, monkeypatch):
    from company_envs.models import strict_schema
    from company_envs.world import grader_author

    root, folder, _ = company
    (root / "config.toml").write_text('[models]\nexpand = ["codex/gpt-test"]\n')
    made = []

    def factory(config, directory):
        made.append((config, directory))
        return Model(TaskGrader(checks=[state_check(0, "$.tickets[0].status")]))

    monkeypatch.setattr(grader_author, "Models", factory)
    draft = author_grader(root, folder, TASK)
    assert draft["receipt"]["job"] == "task_grader"
    assert made[0][0]["models"]["task_grader"] == ["codex/gpt-test"]
    schema = strict_schema(TaskGrader)
    assert schema["additionalProperties"] is False
    assert schema["$defs"]["Predicate"]["additionalProperties"] is False


def test_worker_requirements_cannot_change_after_calibration(company):
    folder, clients, _ = prepare(company, workers=["w1", "w2", "w3"])
    certify(folder, clients)
    path = folder / "tasks" / TASK / "workflow.json"
    workflow = read(path)
    workflow["worker_ids"] = ["w1"]
    write(path, workflow)
    with pytest.raises(CalibrationError, match="requirements changed"):
        grade(folder, TASK, clients)


@pytest.mark.parametrize(("operator", "value"), [("exists", None), ("count_gte", 1)])
def test_seeded_record_predicates_cannot_calibrate(company, operator, value):
    # The value check on 1310 carries this grader over the mechanical floor, so the seeded-record
    # predicate is refused for the reason under test rather than for pinning nothing.
    folder, clients, _ = prepare(
        company,
        [state_check(0, "$.tickets[?(@.id==1310)].status"), state_check(1, "$.tickets", operator, value)],
    )
    with pytest.raises(CalibrationError, match="initial must fail"):
        certify(folder, clients)


def test_mechanical_floor_refuses_a_grader_that_pins_nothing(company):
    """Across nine authored graders, 30 of the 31 state checks used `changed`, which is literally
    `not _equal(before, after)`: setting 440 Zendesk comment bodies to "zzz" and touching nothing
    else passed 8 of 8 state checks. A draft whose checks only prove a field moved goes back to
    the author, and no grader.json is written for it.
    """
    from company_envs.world.grader import MECHANICAL_FLOOR

    _, folder, _ = company
    with pytest.raises(ValueError, match="mechanical floor"):
        prepare(company, [state_check(0, "$.tickets[?(@.id==1310)].status", "changed")])
    assert not (folder / "tasks" / TASK / "grader.json").is_file()

    _, _, model = prepare(company, [state_check(0, "$.tickets[?(@.id==1310)].status")])
    draft = read(folder / "tasks" / TASK / "grader.json")
    assert draft["mechanical_floor"]["pinned_share"] >= MECHANICAL_FLOOR
    assert draft["mechanical_floor"]["unpinned"] == [] and draft["mechanical_floor"]["changed_only"] == 0
    assert "HIDDEN-SOLUTION" not in model.calls[0][1]  # the demand for a value is not a golden leak


def test_calibration_refuses_what_the_zero_one_pair_cannot_see(company):
    """The two existing calibration points accept a `changed` check and so does vandalism, which
    is how a gibberish agent scored 0.60 on all nine graded tasks. The proof recomputes the floor
    from the grader on disk, so a hand-edited draft cannot carry a stale measurement past it.
    """
    from company_envs.world.grader import MECHANICAL_FLOOR

    folder, clients, _ = prepare(company)
    path = folder / "tasks" / TASK / "grader.json"
    draft = read(path)
    draft["grader"]["checks"][0]["predicate"] = {
        "app_id": APP,
        "selector": "$.tickets[?(@.id==1310)].status",
        "operator": "changed",
        "value_json": None,
    }
    draft["mechanical_floor"] = {"state_checks": 1, "pinned": 1, "pinned_share": 1.0, "unpinned": []}
    write(path, draft)
    with pytest.raises(CalibrationError, match="mechanical floor"):
        certify(folder, clients)
    proof = read(folder / "runtime" / "grades" / TASK / "calibration.json")
    # Both older points hold on the very grader the floor rejects: that is what they cannot see.
    assert proof["initial_score"] == 0 and proof["reference_score"] == 1 and not proof["accepted"]
    assert proof["mechanical_floor"]["pinned_share"] == 0 < MECHANICAL_FLOOR
    with pytest.raises(CalibrationError):
        grade(folder, TASK, clients)

    vandalised = deepcopy(clients[APP].initial)
    vandalised["tickets"][0]["status"] = "zzz"
    moved = Predicate(app_id=APP, selector="$.tickets[?(@.id==1310)].status", operator="changed")
    assert evaluate_predicate(moved, {APP: clients[APP].initial}, {APP: vandalised})
    assert not evaluate_predicate(
        predicate("$.tickets[?(@.id==1310)].status"), {APP: clients[APP].initial}, {APP: vandalised}
    )

    folder, clients, _ = prepare(company)
    accepted = certify(folder, clients)
    assert accepted["accepted"] and accepted["mechanical_floor"]["pinned_share"] == 1
    clients[APP].current = vandalised
    assert grade(folder, TASK, clients)["score"] == 0


def test_an_accepted_proof_from_before_the_floor_is_stale(company):
    """Nine graders already carry an accepted calibration.json that predates the floor; four of
    them would still have graded. `grade` recomputes the floor instead of trusting the report,
    so a proof that certifies a grader crediting any edit is refused like any other stale one.
    """
    from company_envs.world.grader import _proof_key

    folder, clients, _ = prepare(company)
    certify(folder, clients)
    clients[APP].current = golden(clients)
    assert grade(folder, TASK, clients)["score"] == 1

    path = folder / "tasks" / TASK / "grader.json"
    draft = read(path)
    draft["grader"]["checks"][0]["predicate"] = {
        "app_id": APP,
        "selector": "$.tickets[?(@.id==1310)].status",
        "operator": "changed",
        "value_json": None,
    }
    write(path, draft)
    proof = read(folder / "runtime" / "grades" / TASK / "calibration.json")
    proof["proof_key"] = _proof_key(draft, {APP: clients[APP].initial})  # as if it were certified
    write(folder / "runtime" / "grades" / TASK / "calibration.json", proof)
    with pytest.raises(CalibrationError, match="stale"):
        grade(folder, TASK, clients)


def test_mechanical_floor_weighs_criteria_the_way_the_score_does():
    """One audited grader hung 8 state checks on a single criterion. Counted per check, one pin
    out of eight would read as 0.125 either way; but the score those checks feed gives the
    criterion one eighth per check, so the floor is weighted per criterion too. `count_gte` and
    `file_exists` are not pins: an edit that appends a record clears them without producing a
    value.
    """
    from company_envs.world.grader import floor_problem, mechanical_floor

    def check(name, ref, operator="changed", value=None):
        return Check(
            id=name,
            criterion_ref=ref,
            kind="state",
            description="Target case updated",
            predicate=predicate("$.tickets[?(@.id==1310)].status", operator, value),
        )

    crowded = [check(f"s{i}", "/success_criteria/0") for i in range(7)]
    crowded.append(check("s7", "/success_criteria/0", "equals", "solved"))
    floor = mechanical_floor(TaskGrader(checks=crowded))
    assert floor["state_checks"] == 8 and floor["pinned"] == 1 and floor["changed_only"] == 7
    assert floor["pinned_share"] == 0.125 and "s0" in floor["unpinned"]
    assert "mechanical floor" in floor_problem(floor)

    split = [crowded[7], check("s0", "/success_criteria/1")]
    floor = mechanical_floor(TaskGrader(checks=split))
    assert floor["pinned_share"] == 0.5 and floor_problem(floor) is None

    counted = [crowded[7], check("s0", "/success_criteria/1", "count_gte", 4)]
    floor = mechanical_floor(TaskGrader(checks=counted))
    assert floor["counted"] == 1 and floor["pinned"] == 1 and floor["pinned_share"] == 0.5
    assert floor_problem(floor) is None  # a quantity beside a value still clears the floor
    only_counts = [check("s0", "/success_criteria/0", "count_gte", 4)]
    assert mechanical_floor(TaskGrader(checks=only_counts))["pinned_share"] == 0


def test_proxy_identity_change_cannot_be_a_state_reward(company):
    _, folder, _ = company
    apps = read(folder / "apps.json")
    apps["apps"][0]["identity_key"] = "currentUser"
    write(folder / "apps.json", apps)
    with pytest.raises(ValueError, match="proxy identity"):
        prepare(company, [state_check(0, "$.currentUser", "changed")])


def test_contributions_identify_proxies_not_network_callers(company):
    """Runtime limitation: one HTTP caller can impersonate multiple contributors.

    This characterizes the missing network boundary for the runtime owner; the
    existing log format cannot distinguish these writes from independent VMs.
    """
    from contextlib import ExitStack
    from urllib.request import Request, urlopen

    from test_hub_app import StubHub

    from company_envs.world.hub_app import HubClient
    from company_envs.world.hub_identity import WorkerAppProxy

    folder, clients, _ = prepare(company, workers=["w1", "w2", "w3"])
    certify(folder, clients)
    initial = clients[APP].initial
    with ExitStack() as cleanup:
        hub = StubHub()
        cleanup.callback(hub.close)
        client = HubClient(hub.url)
        client.seed(SID, initial)
        for worker, status in (("w1", "pending"), ("w2", "open"), ("w3", "solved")):
            proxy = WorkerAppProxy(
                hub.url,
                worker_id=worker,
                sid=SID,
                host="127.0.0.1",
                attribution_log=folder / "runtime/attribution" / f"{APP}.jsonl",
            ).start()
            cleanup.callback(proxy.stop)
            state = deepcopy(initial)
            state["tickets"][0]["status"] = status
            request = Request(
                f"http://127.0.0.1:{proxy.port}/post?sid={SID}",
                data=json.dumps({"action": "set_current", "state": state}).encode(),
                headers={"Content-Type": "application/json"},
            )
            with urlopen(request, timeout=5) as response:
                assert response.status == 200
        result = grade(folder, TASK, {APP: client})
    assert result["contributions"]["observed_workers"] == 3
    assert result["passed"]  # Proxy labels are not authenticated caller identities.
    assert "authenticate callers" in result["contributions"]["basis"]


def test_authoring_view_is_bounded_but_keeps_named_records():
    from company_envs.world.grader import bound_for_authoring

    emails = [{"id": f"e{i}", "subject": f"s{i}", "body": "x" * 3000} for i in range(400)]
    state = {"emails": emails, "labels": [{"id": "l1"}]}
    view = bound_for_authoring(
        state, "gmail_mock", 'initial_materials: ["gmail_mock.emails#e399"]', limit_per_app=100_000
    )
    assert len(json.dumps(view)) < 130_000
    assert view["labels"] == [{"id": "l1"}]
    assert any(r.get("id") == "e399" for r in view["emails"])
    assert any("_sampled" in r for r in view["emails"])


def test_grader_baseline_hash_covers_the_whole_world_not_the_authoring_view(company, monkeypatch):
    import company_envs.world.grader_author as g

    root, folder, clients = company
    state = read(folder / "apps.json")["apps"][0]["state_file"]
    full = read(folder / state)
    full["tickets"] = full["tickets"] + [
        {**full["tickets"][0], "id": 9000 + i, "subject": f"Filler ticket {i}"} for i in range(400)
    ]
    write(folder / state, full)
    monkeypatch.setattr(
        g, "bound_for_authoring", lambda s, app_id, text, limit_per_app=80_000: {"tickets": s["tickets"][:2]}
    )
    prepare((root, folder, clients))
    draft = read(folder / "tasks" / TASK / "grader.json")
    assert draft["initial_hash"] == digest({APP: read(folder / state)})


def file_check(index=0, *, worker="w1", operator="file_contains", **fields):
    return Check(
        id=f"f{index}",
        criterion_ref=f"/success_criteria/{index}",
        kind="state",
        description="Save the delivery decision in Documents",
        predicate=Predicate(operator=operator, path=f"{worker}/Documents/decision.txt", **fields),
    )


def test_file_deliverable_measurement_with_golden_episode(company, tmp_path, capsys):
    """An actual harness action writes the reference file through a guest double."""
    import asyncio
    import shutil

    from test_guest_export import TarTransport, archive_bytes, backend, export

    from company_envs.world.harness import Action, Budget, Episode, FakeBackend, WorkerAgent

    folder, clients, _ = prepare(company)
    certify(folder, clients)
    before_noop = grade(folder, TASK, clients)["score"]
    guest = tmp_path / "guest/Documents"
    guest.mkdir(parents=True)

    class Guest(FakeBackend):
        async def execute(self, action, *, deadline):
            assert action.name == "bash"
            (guest / "decision.txt").write_text("Approve delivery on September 8.")
            return {"exit_code": 0}

    actions = iter([Action("bash", {"command": "write delivery decision"}), Action("done")])
    result = asyncio.run(
        Episode(
            [WorkerAgent("w1", "Planner", lambda obs: next(actions), Guest(), Budget(2, 5))],
            boss_id="w1",
            brief="Save the delivery decision",
            runtime=tmp_path / "reference-episode",
            seconds=5,
        ).run()
    )
    assert result["reason"] == "all_done"
    reference = tmp_path / "reference"
    export(
        backend(
            tmp_path,
            TarTransport(archive_bytes([("Documents/decision.txt", (guest / "decision.txt").read_bytes())])),
            worker="w1",
        ),
        reference / "w1",
    )
    shutil.copytree(reference, folder / "runtime/exports" / TASK)
    before_golden = grade(folder, TASK, clients)["score"]
    shutil.rmtree(folder / "runtime/exports" / TASK)
    folder, clients, model = prepare(company, [file_check(text="Approve delivery on September 8.")])
    assert "sheet_cell" in model.calls[0][1] and "runtime/exports/<task>" in model.calls[0][1]
    proof = calibrate(folder, TASK, clients, reference_diff={}, reference_exports=reference)
    assert (proof["initial_score"], proof["reference_score"]) == (0, 1)
    assert proof["reference_file_hashes"]["w1/Documents/decision.txt"]
    after_noop = grade(folder, TASK, clients)["score"]
    shutil.copytree(reference, folder / "runtime/exports" / TASK)
    after_golden = grade(folder, TASK, clients)["score"]
    assert (before_noop, before_golden, after_noop, after_golden) == (0, 0, 0, 1)
    print(
        f"File fixture: before noop={before_noop}, golden={before_golden}; "
        f"after noop={after_noop}, golden={after_golden}"
    )
    assert "after noop=0.0, golden=1.0" in capsys.readouterr().out


@pytest.mark.parametrize(
    "operator,fields,name",
    [
        ("file_exists", {}, "note.txt"),
        ("file_contains", {"text": "Approved"}, "note.txt"),
        ("sheet_cell", {"sheet": "Decision", "cell": "B2", "value_json": "42"}, "result.xlsx"),
        ("pdf_contains", {"text": "Approved"}, "result.pdf"),
        ("docx_contains", {"text": "Approved"}, "result.docx"),
    ],
)
def test_file_formats_match_content_and_reject_missing_wrong_targets(tmp_path, operator, fields, name):
    from docx import Document
    from fpdf import FPDF
    from openpyxl import Workbook

    documents = tmp_path / "w1/Documents"
    documents.mkdir(parents=True)
    (documents / "note.txt").write_text("Approved")
    book = Workbook()
    book.active.title = "Decision"
    book.active["B2"] = 42
    book.save(documents / "result.xlsx")
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="Approved")
    pdf.output(documents / "result.pdf")
    doc = Document()
    doc.add_paragraph("Decision")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Approved"
    doc.save(documents / "result.docx")
    path = f"w1/Documents/{name}"
    predicate = Predicate(operator=operator, path=path, **fields)
    assert not evaluate_predicate(predicate, {}, {})
    assert evaluate_predicate(predicate, {}, {}, exports=tmp_path)
    wrong = predicate.model_copy(update={"path": f"w2/Documents/{name}"})
    assert not evaluate_predicate(wrong, {}, {}, exports=tmp_path)
    if fields:
        wrong = predicate.model_copy(
            update={"text": "Declined"} if "text" in fields else {"value_json": '"42"'}
        )
        assert not evaluate_predicate(wrong, {}, {}, exports=tmp_path)
    if operator != "file_exists":
        (documents / name).write_bytes(b"\xffbroken")
        with pytest.raises(ValueError, match="cannot read exported file"):
            evaluate_predicate(predicate, {}, {}, exports=tmp_path)


@pytest.mark.parametrize(
    "fields",
    [
        {"path": "../outside"},
        {"path": "/home/ga/Documents/x"},
        {"path": "w1/Documents/../x"},
        {"path": "*/Documents/x"},
        {"path": "w1/private/x"},
        {"app_id": APP},
        {"selector": "$.tickets"},
        {"text": ""},
        {"cell": "A1"},
        {"value_json": '"yes"'},
        {"operator": "sheet_cell", "text": None, "sheet": "Sheet1", "cell": "A0", "value_json": "1"},
        {"operator": "sheet_cell", "text": None, "sheet": "Sheet1", "cell": "A1", "value_json": "[]"},
        {"operator": "pdf_contains", "path": "w1/Documents/*.pdf"},
    ],
)
def test_invalid_file_predicates(fields):
    with pytest.raises(ValueError):
        Predicate(**({"operator": "file_contains", "path": "w1/Documents/result.txt", "text": "ok"} | fields))


def test_globs_skip_symlinks_and_directories(tmp_path):
    directory = tmp_path / "w1/Desktop"
    directory.mkdir(parents=True)
    (directory / "sub").mkdir()
    (directory / "sub/note.txt").write_text("Ready")
    (directory / "link.txt").symlink_to(directory / "sub/note.txt")
    (directory / "shortcut").symlink_to(directory / "sub", target_is_directory=True)
    predicate = Predicate(operator="file_contains", path="w1/Desktop/*.txt", text="Ready")
    assert evaluate_predicate(predicate, {}, {}, exports=tmp_path)
    assert not evaluate_predicate(
        predicate.model_copy(update={"path": "w1/Desktop/link.txt"}), {}, {}, exports=tmp_path
    )
    assert not evaluate_predicate(
        predicate.model_copy(update={"path": "w1/Desktop/shortcut/note.txt"}), {}, {}, exports=tmp_path
    )
    assert not evaluate_predicate(
        Predicate(operator="file_exists", path="w1/Desktop/sub"), {}, {}, exports=tmp_path
    )


def test_initial_file_checks_reject_preexisting_work_and_pin_baseline(company, tmp_path):
    import shutil

    folder, clients, _ = prepare(company, [file_check(text="Approved")])
    reference = tmp_path / "reference"
    target = reference / "w1/Documents/decision.txt"
    target.parent.mkdir(parents=True)
    target.write_text("Approved")
    baseline = folder / "runtime/vms/w1/guest/Documents/decision.txt"
    baseline.parent.mkdir(parents=True)
    baseline.write_text("Approved")
    with pytest.raises(CalibrationError, match="initial must fail"):
        calibrate(folder, TASK, clients, reference_diff={}, reference_exports=reference)
    baseline.write_text("Pending")
    calibrate(folder, TASK, clients, reference_diff={}, reference_exports=reference)
    shutil.copytree(reference, folder / "runtime/exports" / TASK)
    assert grade(folder, TASK, clients)["score"] == 1
    baseline.write_text("Changed baseline")
    with pytest.raises(CalibrationError, match="stale"):
        grade(folder, TASK, clients)


def test_reference_json_files_guards_and_three_worker_credit(company):
    import shutil

    checks = [file_check(i, worker=f"w{i + 1}", text="Approved") for i in range(3)]
    checks[0].guards = [Predicate(operator="file_exists", path="w1/Documents/support.csv")]
    folder, clients, _ = prepare(company, checks, workers=["w1", "w2", "w3"])
    reference = folder / "reference-files"
    for worker in ("w1", "w2", "w3"):
        target = reference / worker / "Documents/decision.txt"
        target.parent.mkdir(parents=True)
        target.write_text("Approved")
    (reference / "w1/Documents/support.csv").write_text("amount\n42")
    write(folder / "tasks" / TASK / "reference.json", {"diff": {}, "exports": "reference-files"})
    assert calibrate(folder, TASK, clients)["accepted"]
    exports = folder / "runtime/exports" / TASK
    shutil.copytree(reference, exports)
    report = grade(folder, TASK, clients)
    assert report["score"] == 1 and report["contributions"]["observed_workers"] == 3
    assert read(folder / "runtime/grades" / TASK / "evidence.json")["exported_files"]
    (exports / "w1/Documents/support.csv").unlink()
    assert grade(folder, TASK, clients)["score"] == pytest.approx(4 / 9)
    with pytest.raises(CalibrationError, match="separate"):
        calibrate(folder, TASK, clients, reference_diff={}, reference_exports=exports)
    with pytest.raises(CalibrationError, match="reference must pass"):
        calibrate(folder, TASK, clients, reference_diff={})


def test_passing_glob_does_not_attribute_unreadable_files(company, tmp_path):
    import shutil

    check = file_check(text="Approved")
    check.predicate = Predicate(operator="file_contains", path="w1/Documents/*.txt", text="Approved")
    folder, clients, _ = prepare(company, [check])
    reference = tmp_path / "reference"
    documents = reference / "w1/Documents"
    documents.mkdir(parents=True)
    (documents / "a.txt").write_text("Approved")
    (documents / "z.txt").write_bytes(b"\xff")
    calibrate(folder, TASK, clients, reference_diff={}, reference_exports=reference)
    shutil.copytree(reference, folder / "runtime/exports" / TASK)
    report = grade(folder, TASK, clients)
    assert report["score"] == 1
    assert report["contributions"]["workers"][0]["changes"] == [
        {"path": "w1/Documents/a.txt", "check_ids": ["f0"]}
    ]


def test_sheet_cell_reads_formula_text_without_recalculation(tmp_path):
    from openpyxl import Workbook

    target = tmp_path / "w1/Documents/budget.xlsx"
    target.parent.mkdir(parents=True)
    book = Workbook()
    book.active["A1"] = "=1+2"
    book.save(target)
    predicate = Predicate(
        operator="sheet_cell", path="w1/Documents/budget.xlsx", sheet="Sheet", cell="A1", value_json='"=1+2"'
    )
    assert evaluate_predicate(predicate, {}, {}, exports=tmp_path)
    assert not evaluate_predicate(predicate.model_copy(update={"value_json": "3"}), {}, {}, exports=tmp_path)
    with pytest.raises(ValueError, match="cannot read exported file"):
        evaluate_predicate(predicate.model_copy(update={"sheet": "Missing"}), {}, {}, exports=tmp_path)


def test_world_material_is_initial_desktop_evidence_before_launch(company, tmp_path):
    _, folder, clients = company
    material = folder / "world/materials/w1/decision.txt"
    material.parent.mkdir(parents=True)
    material.write_text("Approved")
    check = file_check(text="Approved")
    check.predicate.path = "w1/Desktop/decision.txt"
    prepare(company, [check])
    reference = tmp_path / "reference/w1/Desktop/decision.txt"
    reference.parent.mkdir(parents=True)
    reference.write_text("Approved")
    with pytest.raises(CalibrationError, match="initial must fail"):
        calibrate(folder, TASK, clients, reference_diff={}, reference_exports=tmp_path / "reference")


def test_text_glob_can_find_valid_file_after_binary_file(tmp_path):
    documents = tmp_path / "w1/Documents"
    documents.mkdir(parents=True)
    (documents / "a.bin").write_bytes(b"\xff")
    (documents / "z.txt").write_text("Approved")
    predicate = Predicate(operator="file_contains", path="w1/Documents/*", text="Approved")
    assert evaluate_predicate(predicate, {}, {}, exports=tmp_path)


def test_golden_app_replay_accepts_separate_file_reference(company, tmp_path):
    class ReplayClient:
        def __init__(self):
            self.states = {}

        def current(self, sid):
            return {"has_custom_state": sid in self.states}

        def seed(self, sid, state):
            self.states[sid] = deepcopy(state)

        def inspect(self, sid):
            return {"initial_state": self.states[sid], "current_state": self.states[sid], "state_diff": {}}

    checks = [file_check(i, worker=f"w{i + 1}", text="Approved") for i in range(3)]
    folder, _, _ = prepare(company, checks, workers=["w1", "w2", "w3"])
    reference = tmp_path / "reference"
    for worker in ("w1", "w2", "w3"):
        path = reference / worker / "Documents/decision.txt"
        path.parent.mkdir(parents=True)
        path.write_text("Approved")
    write(
        folder / "tasks" / TASK / "golden.json",
        [{"worker_id": "w1", "app_id": APP, "action": "message", "text": "Files were exported separately."}],
    )
    proof = calibrate(folder, TASK, {APP: ReplayClient()}, golden=True, reference_exports=reference)
    assert proof["accepted"] and proof["reference_score"] == 1
    assert proof["contributions"]["observed_workers"] == 3
    assert proof["golden_sid"].startswith("golden-")


def test_judgment_packet_is_bounded_under_the_model_input_limit():
    from company_envs.world.grader import _bounded

    big = {"id": "r1", "body": "x" * 50_000}
    records = {f"r{i}": dict(big, id=f"r{i}") for i in range(40)}
    packet = {
        "checks": [
            {
                "check": {"id": "c"},
                "evidence": [
                    {
                        "id": "app:0",
                        "current": {
                            "changed_records": records,
                            "unchanged_index": [{"id": str(i), "label": "l"} for i in range(5000)],
                            "unchanged_count": 5000,
                        },
                    }
                ],
            }
        ]
    }
    import json

    assert len(json.dumps(packet)) > 900_000
    bounded = _bounded(packet)
    assert len(json.dumps(bounded)) <= 900_000
    current = bounded["checks"][0]["evidence"][0]["current"]
    assert current["index_truncated"] and len(current["unchanged_index"]) == 100
    assert all(len(r["body"]) < 2000 for r in current["changed_records"].values())
    assert set(current["changed_records"]) == set(records)  # every changed record survives


def test_compact_unwraps_selector_matches_and_bound_has_a_hard_fallback():
    import json

    from company_envs.world.grader import _bounded, _compact

    emails = [{"id": f"m{i}", "subject": "s", "body": "b" * 100} for i in range(300)]
    after = emails[:-1] + [dict(emails[-1], body="edited")] + [{"id": "new", "subject": "n", "body": "x"}]
    compact = _compact([emails], [after])  # select() wraps the collection in a match list
    assert set(compact["changed_records"]) == {"m299", "new"} and compact["unchanged_count"] == 299

    raw = {
        "checks": [
            {
                "check": {"id": "c"},
                "evidence": [
                    {
                        "id": "app:0",
                        "current": [emails * 40],
                        "initial": [emails * 40],
                        "state_diff": emails * 40,
                    }
                ],
            }
        ]
    }
    assert len(json.dumps(raw)) > 900_000
    bounded = _bounded(raw)
    assert len(json.dumps(bounded)) <= 900_000
    assert bounded["checks"][0]["evidence"][0]["current"]["truncated"] is True


def test_full_calibration_judges_every_check_on_the_reference(company):
    """A judged check must pass on the golden and show no change on the untouched initial."""
    judged = Check(
        id="c1",
        criterion_ref="/success_criteria/1",
        kind="judgment",
        description="Assess the resolved ticket",
        rubric="The ticket resolution is substantive.",
        app_paths=[AppPath(app_id=APP, selector="$.tickets[0]")],
    )
    folder, clients, _model = prepare(company, [state_check(0, "$.tickets[0].status"), judged])
    reference = {APP: {"tickets": golden(clients)["tickets"]}}

    class CalibratedModel(Model):
        def call(self, job, prompt, response_type):
            result, receipt = super().call(job, prompt, response_type)
            packet = json.loads(prompt[prompt.index("{") :])
            if not any(e["changed"] for e in packet["checks"][0]["evidence"]):
                result = TaskJudgment(
                    checks=[
                        Judgment(check_id="c1", status="fail", reason="unchanged", evidence_refs=["app:0"])
                    ]
                )
            return result, receipt

    passing = CalibratedModel(
        TaskJudgment(checks=[Judgment(check_id="c1", status="pass", reason="ok", evidence_refs=["app:0"])])
    )
    report = calibrate(folder, TASK, clients, reference_diff=reference, models=passing)
    assert report["accepted"] and report["scope"] == "all_checks"
    assert [(r["id"], r["reference"]) for r in report["judged_checks"]] == [("c1", "pass")]
    assert passing.calls[-1][0] == "task_judgment"

    indiscriminate = Model(
        TaskJudgment(checks=[Judgment(check_id="c1", status="pass", reason="ok", evidence_refs=["app:0"])])
    )
    with pytest.raises(CalibrationError, match="falsely passed the untouched initial"):
        calibrate(folder, TASK, clients, reference_diff=reference, models=indiscriminate)
    assert not read(folder / "runtime/grades" / TASK / "calibration.json")["accepted"]

    failing = Model(
        TaskJudgment(checks=[Judgment(check_id="c1", status="fail", reason="thin", evidence_refs=["app:0"])])
    )
    with pytest.raises(CalibrationError, match="judge must pass the reference"):
        calibrate(folder, TASK, clients, reference_diff=reference, models=failing)

    partial = calibrate(folder, TASK, clients, reference_diff=reference)
    assert partial["scope"] == "state_checks_only"  # no judge: partial, and the agreement rejects it


def test_export_evidence_reads_deliverables_and_marks_changes(tmp_path):
    from openpyxl import Workbook

    from company_envs.world.grader import _export_evidence, _file_text, mechanical_floor

    before = tmp_path / "before" / "w" / "Documents"
    after = tmp_path / "after" / "w" / "Documents"
    before.mkdir(parents=True)
    after.mkdir(parents=True)
    (before / "plan.txt").write_text("draft")
    (after / "plan.txt").write_text("final plan: nursing Tuesday 09:00")
    book = Workbook()
    book.active.title = "Costs"
    book.active["A1"], book.active["B1"] = "Item", 510
    book.save(after / "costs.xlsx")
    (after / "unrelated.txt").write_text("x")
    exports = {
        "w/Documents/plan.txt": after / "plan.txt",
        "w/Documents/costs.xlsx": after / "costs.xlsx",
        "w/Documents/unrelated.txt": after / "unrelated.txt",
    }
    initial = {"w/Documents/plan.txt": before / "plan.txt"}

    check = Check(
        id="a1",
        criterion_ref="/success_criteria/0",
        kind="artifact",
        description="Plan and costs agree",
        rubric="The plan names the visits the cost sheet prices.",
        app_paths=[AppPath(app_id=APP, selector="$.tickets")],
        material_paths=["*plan*", "*.xlsx"],
    )
    items, errors = _export_evidence(check, exports, initial)
    assert errors == [] and [i["path"] for i in items] == ["w/Documents/costs.xlsx", "w/Documents/plan.txt"]
    by_path = {i["path"]: i for i in items}
    assert (
        by_path["w/Documents/plan.txt"]["changed"]
        and "nursing Tuesday" in by_path["w/Documents/plan.txt"]["content"]
    )
    assert (
        "## Costs" in by_path["w/Documents/costs.xlsx"]["content"]
        and "510" in by_path["w/Documents/costs.xlsx"]["content"]
    )
    assert "unrelated" not in by_path
    assert _file_text(after / "plan.txt").startswith("final plan")

    check.material_paths = []
    items, _ = _export_evidence(check, exports, initial)
    assert len(items) == 3  # a check naming no files sees every export, capped

    grader = TaskGrader(checks=[state_check(0, "$.tickets[0].status"), check])
    floor = mechanical_floor(grader)
    assert (
        floor["state_checks"] == 1 and floor["judged"] == 1 and floor["pinned"] + floor["changed_only"] == 1
    )


def test_grader_feedback_never_carries_judge_reasons():
    from company_envs.world.grader import grader_safe_feedback

    message = (
        "grader rejected: EVERY state check must meet calibration; c1: judge must pass the reference "
        "(fail: 3 of 3 votes: the golden writes msg-plastic-withdrawal with $43,200); c2: initial must fail"
    )
    safe = grader_safe_feedback(message)
    assert "msg-plastic" not in safe and "43,200" not in safe
    assert "c1: the judge did not pass the reference outcome" in safe and "c2: initial must fail" in safe


def test_judged_checks_may_not_select_a_whole_collection():
    from company_envs.world.grader import authoring_problems

    whole = Check(
        id="j1",
        criterion_ref="/success_criteria/0",
        kind="judgment",
        description="Report quality",
        rubric="The report is complete.",
        app_paths=[AppPath(app_id=APP, selector="$.tickets")],
    )
    grader = TaskGrader(checks=[state_check(0, "$.tickets[0].status"), whole])
    with pytest.raises(ValueError, match="selects a whole collection"):
        authoring_problems(grader)
    TaskGrader.model_validate(grader.model_dump())  # graders calibrated before this rule still load


def test_wiped_collections_flags_an_app_reset_to_demo_data():
    from company_envs.world.grader import _wiped_collections

    initial = {"docs": {"documents": [{"id": f"d{i}", "title": "seeded"} for i in range(12)], "ui": {}}}
    final = {"docs": {"documents": [{"id": "doc-1", "title": "Project Proposal"}], "ui": {}}}
    assert _wiped_collections(initial, final) == [{"collection": "docs.documents", "seeded": 12, "lost": 12}]
    assert _wiped_collections(initial, initial) == []


def test_every_whole_collection_selector_is_reported_not_just_the_first():
    """A judged check's app_paths must name records, and authoring_problems named one offender at a
    time out of a three-attempt budget. Measured over the 90 grader drafts on disk: 59 were rejected
    here and 55 of those 59 (93%) carried more than one whole-collection selector, a median of ten,
    so an author that fixed the reported one heard about the next one and ran out of attempts."""
    from company_envs.world.grader import authoring_problems

    judged = Check(
        id="j1",
        criterion_ref="/success_criteria/0",
        kind="judgment",
        description="Report quality",
        rubric="The report is complete.",
        app_paths=[
            AppPath(app_id=APP, selector="$.tickets"),
            AppPath(app_id=APP, selector="$.comments"),
            AppPath(app_id=APP, selector="$.tickets[?(@.id==1310)].subject"),
        ],
    )
    grader = TaskGrader(checks=[state_check(1, "$.tickets[?(@.id==1310)].status"), judged])
    with pytest.raises(ValueError, match="selects a whole collection") as caught:
        authoring_problems(grader)
    assert "Fix all 2" in str(caught.value)
    assert "$.tickets" in str(caught.value) and "$.comments" in str(caught.value)
    assert "1310" not in str(caught.value)  # the selector that already names a record is left alone


@pytest.mark.parametrize(
    "problem",
    [
        "j1: Zendesk_mock $.tickets selects a whole collection",
        "c1: proxy identity cannot establish task progress",
        "c1: isDragging is interface state (a sort order, an open view), not task progress",
        "mechanical floor: value checks carry 0.00 of the mechanical score",
    ],
)
def test_every_rejection_a_gate_raises_carries_a_recipe(problem):
    """A rule handed back without an example of obeying it costs one of three attempts. Measured:
    61 of the 90 recorded drafts (68%) were rejected for a whole-collection or proxy-identity
    selector, and neither kind had a hint at all."""
    from company_envs.world.grader_author import _draft_hint

    hint = _draft_hint(problem, ["Zendesk_mock.tickets"])
    assert hint.startswith(" | ") and len(hint) > 120


def test_the_grader_author_is_told_which_collections_are_identity(company):
    """_validate_grader refuses a scored state check on an app's identity collection, and the author
    could not see which collections those were. One company's only task died after six drafts that
    all put "make sure the shared service bulletin is accurate" on the Slack profile, because the
    bulletin is a profile status message and nothing in the payload said that was off limits."""
    root, folder, _ = company
    apps = read(folder / "apps.json")
    apps["apps"][0]["identity_key"] = "currentUser"
    write(folder / "apps.json", apps)
    payload = authoring_payload(root, folder, TASK)
    assert payload["contract"]["identity_collections"] == {APP: "currentUser"}
    # and the instructions say what that means, so the rule can be obeyed rather than discovered.
    from company_envs.world.grader_author import AUTHOR_INSTRUCTIONS

    assert "contract.identity_collections" in AUTHOR_INSTRUCTIONS


@pytest.mark.parametrize(
    "failure,earlier,target",
    [
        ("grader rejected: EVERY state check; c1: initial must fail", (), "grader"),
        ("grader rejected: EVERY state check; c1: judge must pass the reference (fail)", (), "golden"),
        ("grader rejected: EVERY state check; c1: reference must pass", (), "grader"),
        (
            "grader rejected: EVERY state check; c1: reference must pass",
            ("grader rejected: EVERY state check; c1: reference must pass",),
            "stop",
        ),
        (
            "grader rejected: EVERY state check; c2: reference must pass",
            ("grader rejected: EVERY state check; c1: reference must pass",),
            "grader",
        ),
    ],
)
def test_a_reference_that_never_passes_is_not_the_graders_fault(failure, earlier, target):
    """Calibration branched on the words "judge must pass the reference" alone, so a state check
    whose reference never passes was handed to the grader author round after round. Measured:
    nordstrom's service_request_answer and national-archives' six 6101-6106_obligations_updated
    checks reported "reference must pass" in all three attempts, and the judges' reasons named
    world defects ("the supplied Options sheet remains headers-only")."""
    from company_envs.world.grader_author import repair_target

    chosen, reason = repair_target(failure, earlier)
    assert chosen == target and reason


def test_every_contract_problem_in_a_draft_is_reported_together(company):
    """_validate_grader named the first broken rule and stopped. Measured: 8 of the 21 recorded
    drafts that broke the contract broke it in more than one place (up to four), and one company
    spent six drafts alternating between an interface-state selector and a whole-collection one,
    hearing about whichever came first. Both recipes now come back with the message."""
    from company_envs.world.grader import _validate_grader
    from company_envs.world.grader_author import _draft_hint

    _, folder, _ = company
    apps = read(folder / "apps.json")
    apps["apps"][0]["identity_key"] = "currentUser"
    write(folder / "apps.json", apps)
    initial = read(folder / "world" / f"{APP}.state.json")
    initial["currentSortColumn"] = "status"
    grader = TaskGrader(
        checks=[
            state_check(0, "$.currentUser.id", operator="changed"),
            state_check(1, "$.currentSortColumn", operator="changed"),
            state_check(2, "$.tickets[?(@.id==1310)].status"),
        ]
    )
    criteria = [{"requirement": f"r{i}", "method": "state"} for i in range(3)]
    with pytest.raises(ValueError) as caught:
        _validate_grader(grader, criteria, {APP: initial}, folder, [f"{APP}.comments"])
    message = str(caught.value)
    assert "proxy identity cannot establish task progress" in message
    assert "currentSortColumn is interface state" in message
    assert "no check selects decisive collection(s)" in message
    hint = _draft_hint(message, [f"{APP}.comments"])
    assert "identity_collections" in hint and "select records in" in hint


def test_a_hub_that_was_not_serving_is_not_a_verdict_and_spends_no_round():
    """An empty run must not write the receipt a real failure writes. Measured: 2 of the 7 recorded
    calibration failures spent their last round on "golden replay: <urlopen error [Errno 111]
    Connection refused>" and were written up as a verdict, retiring two tasks nothing had been
    measured about. A lost round must also not clear a stuck reference: without that, a dead socket
    between two identical reference failures cancels the stop and buys three more judge votes."""
    from company_envs.world.calibration import CalibrationUnavailable
    from company_envs.world.grader_author import environment_fault, repair_target

    dead = "golden replay: <urlopen error [Errno 111] Connection refused>"
    assert environment_fault(dead)
    assert environment_fault(CalibrationUnavailable(dead))
    assert environment_fault("cannot inspect calibration baseline: {'gmail_mock': '...'}")
    assert not environment_fault("grader rejected: EVERY state check; c1: reference must pass")
    assert repair_target(dead)[0] == "environment"
    stuck = "grader rejected: EVERY state check; c1: reference must pass"
    assert repair_target(stuck, [dead, stuck])[0] == "stop"
    assert repair_target(stuck, [dead])[0] == "grader"  # one real round so far, not two


def test_a_judge_that_could_not_see_the_evidence_reports_not_proven(company):
    """One calibration round was spent on "not visible in the truncated evidence".

    A judgment packet is bounded to fit the model input limit, and the packet holds exactly one
    check -- so any reduction in it is a reduction of the evidence that verdict rests on. A fail on
    a reduced packet is a statement about the packet, not about the world. Scoring does not change
    (a pending judged check earns zero, as a fail does); what changes is that the repair round is no
    longer spent re-authoring a golden over evidence the judge never received.
    """
    from company_envs.world.judge import JUDGE_PROMPT, Judgment, TaskJudgment, _bounded, _judge_one

    bulky = {"changed_records": {str(n): {"body": "x" * 5_000} for n in range(400)}}
    item = {
        "check": {"id": "c1", "kind": "semantic", "description": "The memo names the supplier"},
        "evidence": [{"id": "e1", "changed": True, "current": bulky}],
        "errors": [],
    }
    packet = _bounded({"checks": [item]}, limit=20_000)
    assert packet["evidence_truncated"]
    assert "evidence_truncated" in JUDGE_PROMPT

    class Judge:
        def __init__(self):
            self.prompts = []

        def call(self, job, prompt, response_type):
            self.prompts.append(prompt)
            return (
                TaskJudgment(
                    checks=[
                        Judgment(
                            check_id="c1",
                            status="fail",
                            reason="the memo is not visible in the truncated evidence",
                            evidence_refs=["e1"],
                        )
                    ]
                ),
                {"job": job},
            )

    models = Judge()
    verdict, error, _receipts = _judge_one(models, {}, item)
    assert error is None
    from company_envs.world.judge import TRUNCATED_VERDICT

    assert verdict.status == "pending"
    assert verdict.reason.startswith(TRUNCATED_VERDICT)
    assert "not visible in the truncated evidence" in verdict.reason
    # A fail on evidence that did fit is still a fail.
    small = {
        **item,
        "evidence": [{"id": "e1", "changed": True, "current": {"changed_records": {"1": {"body": "s"}}}}],
    }
    verdict, _error, _receipts = _judge_one(models, {}, small)
    assert verdict.status == "fail"
