"""Cheap canonical and desktop invariants, independent of model review."""

import csv
import io
import re
from collections import Counter
from datetime import date
from email import policy
from email.parser import BytesParser
from pathlib import Path

from .documents import _sheets, render_material


def history_checks(history, spec, reference_date):
    errors, months, ids = [], Counter(), []
    for item in history:
        if not isinstance(item, dict):
            errors.append("history records must be objects")
            continue
        ids.append(item.get("id"))
        try:
            day = date.fromisoformat(item["date"])
            if day > date.fromisoformat(reference_date[:10]):
                errors.append(f"history follows reference date: {item.get('id')}")
            months[str(day)[:7]] += 1
        except (KeyError, TypeError, ValueError):
            errors.append("every history record needs an ISO date")
    valid_ids = [i for i in ids if isinstance(i, str) and i.strip()]
    if len(valid_ids) != len(history) or len(set(valid_ids)) != len(valid_ids):
        errors.append("history records need unique nonempty string IDs")
    if len(history) < spec.get("minimum_history_records", 1) or len(months) < spec.get("history_months", 1):
        errors.append("actual canonical history does not cover the requested months and event count")
    if spec.get("history_start"):
        start = date.fromisoformat(spec["history_start"])
        count = spec.get("history_months", 1)
        first = start.year * 12 + start.month - 1
        expected = [f"{i // 12:04d}-{i % 12 + 1:02d}" for i in range(first, first + count)]
        missing = [m for m in expected if months[m] < spec.get("minimum_history_per_month", 1)]
        if missing:
            errors.append(f"history does not cover the required window: {', '.join(missing)}")
    return errors, {m: n for m, n in sorted(months.items()) if n}


def record_checks(records, world, app_ids):
    errors = []
    known = {r.get("id") for r in records if isinstance(r, dict) and isinstance(r.get("id"), str)}
    # A history event can be the canonical definition of an older transaction.
    known.update(
        r["record_id"]
        for r in world.get("history", [])
        if isinstance(r, dict) and isinstance(r.get("record_id"), str)
    )
    references = {
        "account_id",
        "contact_id",
        "product_id",
        "actor_id",
        "owner_id",
        "campaign_id",
        "vendor_id",
        "order_id",
    }
    for row in records:
        if not isinstance(row, dict):
            continue
        destinations = row.get("apps", [row["app_id"]] if row.get("app_id") else [])
        if (
            not isinstance(destinations, list)
            or not destinations
            or any(a not in app_ids for a in destinations)
        ):
            errors.append(f"{row.get('id')}: missing or unknown app destinations")
        for field in references & row.keys():
            value = row[field]
            if value is not None and (not isinstance(value, str) or value not in known):
                errors.append(f"{row.get('id')}.{field}: unknown canonical record {value!r}")
    return errors


def _normalized(text):
    return re.sub(r"[\W_]+", "", text.casefold())


def material_check(material):
    """Reject trivial/invalid sources and verify the actual rendered document.

    Small notes are valid. Length is only a stub screen, never a claim of semantic
    quality. Rendering/parsing is local and runs before paying for model review.
    """
    text, suffix = material.content.strip(), Path(material.path).suffix.lower()
    label = f"{material.worker_id}/{material.path}"
    try:
        if suffix in {".csv", ".xlsx"}:
            blocks = _sheets(text) if suffix == ".xlsx" else [("CSV", text)]
            for name, body in blocks:
                rows = list(csv.reader(io.StringIO(body.strip()), strict=True))
                if len(rows) < 2 or not any(any(c.strip() for c in r) for r in rows[1:]):
                    raise ValueError(f"{name}: needs a header and actual data")
                if suffix == ".csv" and any(len(r) != len(rows[0]) for r in rows):
                    raise ValueError("CSV rows disagree with the header")
        elif len(_normalized(text)) < 20 or len(set(re.findall(r"\w+", text.casefold()))) < 3:
            raise ValueError("material is empty or a trivial stub")
        data = render_material(material.path, material.content.encode())
        stream = io.BytesIO(data)
        rendered = None
        if suffix == ".pdf":
            from pypdf import PdfReader

            rendered = "".join(p.extract_text() or "" for p in PdfReader(stream).pages)
        elif suffix == ".docx":
            from docx import Document

            rendered = "".join(p.text for p in Document(stream).paragraphs)
        elif suffix == ".pptx":
            from pptx import Presentation

            rendered = "".join(
                s.text for slide in Presentation(stream).slides for s in slide.shapes if s.has_text_frame
            )
        elif suffix == ".xlsx":
            from openpyxl import load_workbook
            from openpyxl.formula.tokenizer import Tokenizer

            book = load_workbook(stream)
            for sheet in book:
                for row in sheet:
                    for cell in row:
                        if cell.data_type == "e":
                            raise ValueError(
                                f"{sheet.title}!{cell.coordinate}: spreadsheet error {cell.value}"
                            )
                        if cell.data_type == "f":
                            tokens = Tokenizer(cell.value).items
                            if not tokens or any(t.subtype == "ERROR" for t in tokens):
                                raise ValueError(f"{sheet.title}!{cell.coordinate}: invalid formula")
        elif suffix == ".eml":
            mail = BytesParser(policy=policy.default).parsebytes(data)
            if (
                mail.defects
                or not all(mail.get(k) for k in ("From", "To", "Subject"))
                or mail.get_body() is None
            ):
                raise ValueError("email needs valid headers and a body")
        if rendered is not None and _normalized(rendered) != _normalized(text):
            raise ValueError("rendered document lost or changed source text")
    except Exception as exc:  # noqa: BLE001 - parser failures are reported as failed material checks
        return f"{label}: {type(exc).__name__}: {exc}"
    return None
