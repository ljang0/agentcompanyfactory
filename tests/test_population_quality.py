from copy import deepcopy

import pytest

from company_envs.storage import read
from company_envs.world.population_quality import (
    materialize_drive,
    normalize_slack,
    quality_errors,
    resource_errors,
    sync_documents,
)


def document():
    return {
        "id": "d",
        "title": "Working note",
        "content": "<p>Confirm the quantity.</p>",
        "ownerId": "owner",
        "created": "2026-01-01T10:00:00Z",
        "updated": "2026-01-02T10:00:00Z",
        "sharedWith": [{"userId": "reader", "permission": "viewer"}],
    }


def test_won_deals_cannot_retain_loss_reasons():
    state = {
        "hubspot_mock": {
            "deals": [
                {"id": "order", "stage": "closed_won", "closedLostReason": "Customer abandoned project"}
            ]
        }
    }
    assert quality_errors(state) == ["deal order: won business retains a loss reason"]
    state["hubspot_mock"]["deals"][0]["closedLostReason"] = ""
    assert not quality_errors(state)


def test_slack_uses_rendered_record_as_only_source():
    state = {
        "messages": {
            "ch": [
                {"messageId": "p", "threadId": "t"},
                {"messageId": "r", "threadId": "t", "content": "Actual message", "senderId": "a"},
            ]
        },
        "threads": {
            "t": {
                "threadId": "t",
                "channelId": "ch",
                "parentMessageId": "p",
                "replies": [{"messageId": "r", "threadId": "t", "content": "Contradiction", "senderId": "b"}],
            }
        },
    }
    assert normalize_slack(state)["conflicting_inline_copies_removed"] == ["r"]
    assert state["threads"]["t"]["replies"] == ["r"]
    assert state["messages"]["ch"][1]["content"] == "Actual message"
    assert state["messages"]["ch"][0]["threadId"] is None
    before = deepcopy(state)
    normalize_slack(state)
    assert state == before
    assert quality_errors({"slack_mock": state}) == []


def test_slack_native_ids_drive_reply_notification_and_channel_chronology():
    slack = {
        "messages": {
            "ch": [
                {"messageId": "p", "timestamp": "2026-05-01T09:00:00Z"},
                {"messageId": "r", "threadId": "t", "timestamp": "2026-04-30T09:00:00Z"},
            ]
        },
        "threads": {"t": {"threadId": "t", "parentMessageId": "p", "replies": ["r"]}},
        "notifications": [{"notificationId": "n", "messageId": "p", "timestamp": "2026-04-30T09:00:00Z"}],
        "channels": [{"channelId": "ch", "createdAt": "2026-05-01T09:00:00Z"}],
    }
    errors = quality_errors({"slack_mock": slack})
    assert len(errors) == 3
    slack["messages"]["ch"][1]["timestamp"] = "2026-05-01T10:00:00Z"
    slack["notifications"][0]["timestamp"] = "2026-05-01T09:00:00Z"
    assert not quality_errors({"slack_mock": slack})


def test_document_projection_preserves_private_access_and_detects_drift():
    docs, drive = {"documents": {"d": document()}, "comments": []}, {"items": {}}
    sync_documents(docs, drive)
    assert drive["items"]["d"]["parentId"] is None
    states = {"google_docs_mock": docs, "google_drive_mock": drive}
    assert quality_errors(states) == []
    drive["items"]["d"]["sharedWith"][0]["role"] = "editor"
    assert any("access differs" in e for e in quality_errors(states))
    drive["items"]["d"]["mimeType"] = "application/pdf"
    with pytest.raises(ValueError, match="explicit source"):
        sync_documents(docs, drive)


def test_comments_require_quote_parent_date_and_permission():
    docs = {
        "documents": {"d": document()},
        "comments": [
            {
                "id": "c",
                "docId": "d",
                "userId": "reader",
                "created": "2025-01-01T10:00:00Z",
                "quotedText": "An invented fact",
                "content": "Check it.",
                "replies": [],
            }
        ],
    }
    errors = quality_errors({"google_docs_mock": docs})
    assert len(errors) == 3


def test_office_downloads_are_parseable_and_bound_to_resources(tmp_path):
    import base64
    import io

    from docx import Document
    from openpyxl import load_workbook
    from pypdf import PdfReader

    drive = {
        "items": {
            ext: {"id": ext, "name": "report" + ext, "size": 0, "content": "Quantity,Total\n12,48"}
            for ext in [".pdf", ".xlsx", ".docx"]
        }
    }
    materialize_drive(tmp_path, drive)
    binary = {ext: base64.b64decode(r["thumbnailUrl"].split(",", 1)[1]) for ext, r in drive["items"].items()}
    assert PdfReader(io.BytesIO(binary[".pdf"])).pages
    assert load_workbook(io.BytesIO(binary[".xlsx"])).active.max_row == 2
    assert Document(io.BytesIO(binary[".docx"])).paragraphs
    assert quality_errors({"google_drive_mock": drive}) == []
    assert resource_errors(tmp_path) == []
    manifest = read(tmp_path / "world/population/files/RESOURCES.json")
    (tmp_path / manifest[0]["path"]).write_bytes(b"bad")
    assert resource_errors(tmp_path)
